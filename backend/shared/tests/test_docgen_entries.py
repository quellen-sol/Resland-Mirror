"""Entry-line and header capture/render: per-section title order and
separator, head/tail styles that follow the source's positions, date
ranges without a start, stacked/labelled contact blocks, and loading
profiles written by older or newer capturers.
"""

import io
import json
import unittest
from dataclasses import asdict, dataclass, field

import docx
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Inches
from shared.docgen import (
    CONTACT_INLINE,
    CONTACT_STACKED,
    ROLE_DATES,
    ROLE_EDUCATION_SUB,
    ROLE_ENTRY_LOCATION,
    ROLE_ENTRY_SUB,
    ROLE_ENTRY_TITLE,
    FormatProfile,
    ResumeContent,
    ResumeEducation,
    ResumeExperience,
    ResumeHeader,
    ResumeProject,
    _dates,
    _skills_category_split,
    _title_separator,
    capture_format_profile,
    capture_format_profile_pdf,
    content_source_hash,
    default_format_profile,
    render_docx,
)
from shared.pdfdoc import PdfDoc, PdfLine, PdfPage, PdfRun, PdfStyle


@dataclass
class _Basic:
    name: str = ""
    email: str = ""
    phone: str = ""
    location: str = ""
    links: list = field(default_factory=list)


@dataclass
class _Bullet:
    text: str


@dataclass
class _Exp:
    company: str
    position: str
    start_date: str
    end_date: str | None
    bullets: list = field(default_factory=list)
    location: str = ""


@dataclass
class _Edu:
    institution: str
    degree: str
    field_of_study: str = ""
    start_date: str = ""
    end_date: str | None = None
    description: str = ""


@dataclass
class _Proj:
    name: str
    link: str | None
    description: str
    start_date: str
    end_date: str | None


@dataclass
class _Profile:
    summary: str | None
    basic_info: _Basic
    experience: list = field(default_factory=list)
    skills: list = field(default_factory=list)
    education: list = field(default_factory=list)
    projects: list = field(default_factory=list)


def _profile() -> _Profile:
    return _Profile(
        summary=None,
        basic_info=_Basic(
            name="Jane Doe", email="jane@x.com", phone="555-1234", location="Denver, CO"
        ),
        experience=[
            _Exp("Acme Corp", "Software Engineer", "Jan 2020", None, [_Bullet("Built the thing")])
        ],
        education=[_Edu("State University", "BS", start_date="2016", end_date="2020")],
        projects=[
            _Proj("Widgetizer", "github.com/jane/widgetizer", "A widget tool", "2021", None)
        ],
    )


def _content() -> ResumeContent:
    return ResumeContent(
        header=ResumeHeader(
            name="Jane Doe", email="jane@x.com", phone="555-1234", location="Denver, CO"
        ),
        experiences=[
            ResumeExperience("Acme Corp", "Software Engineer", "Jan 2020", None, ["Built the thing"])
        ],
        education=[ResumeEducation("State University", "BS", "", "2016", "2020")],
        projects=[
            ResumeProject("Widgetizer", "github.com/jane/widgetizer", "A widget tool", "2021", None)
        ],
    )


def _paras(blob: bytes):
    return [p for p in docx.Document(io.BytesIO(blob)).paragraphs if p.text.strip()]


def _para(paras, needle):
    return next(p for p in paras if needle in p.text)


def _runs(para, *spec):
    """Add runs (text, bold, italic) to a docx paragraph."""
    for text, bold, italic in spec:
        run = para.add_run(text)
        run.font.bold = bold
        run.font.italic = italic


def _company_first_doc() -> Document:
    """Company-first experience (company bold), title-first education with
    a comma, project with a pipe — each section its own convention."""
    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("jane@x.com | 555-1234 | Denver, CO")
    doc.add_paragraph("EXPERIENCE")
    _runs(
        doc.add_paragraph(),
        ("Acme Corp", True, False),
        (" - Software Engineer", False, False),
        ("\tJan 2020 – Present", False, False),
    )
    doc.add_paragraph("• Built the thing")
    doc.add_paragraph("EDUCATION")
    _runs(
        doc.add_paragraph(),
        ("BS", True, False),
        (", State University", False, True),
        ("\t2016 – 2020", False, False),
    )
    doc.add_paragraph("PROJECTS")
    _runs(
        doc.add_paragraph(),
        ("Widgetizer", True, False),
        (" | github.com/jane/widgetizer", False, False),
        ("\t2021", False, False),
    )
    doc.add_paragraph("A widget tool")
    return doc


def _pdf_line(text, top, runs=None, *, size=11.0, bold=False, x0=72.0, gap=4.0):
    style = PdfStyle(font_name="Helvetica", size_pt=size, bold=bold)
    return PdfLine(
        text=text,
        page=0,
        x0=x0,
        x1=540.0,
        top=top,
        bottom=top + size + 2,
        style=style,
        gap_before_pt=gap,
        runs=runs or [],
    )


def _run(text, bold=False, italic=False, color=None, x0=None):
    return PdfRun(
        text=text,
        style=PdfStyle(
            font_name="Helvetica", size_pt=11.0, bold=bold, italic=italic, color=color
        ),
        x0=x0,
    )


def _company_first_pdf() -> PdfDoc:
    """The same resume as pdf IR lines carrying style runs."""
    lines = [
        _pdf_line("Jane Doe", 60, size=22, bold=True),
        _pdf_line("jane@x.com | 555-1234 | Denver, CO", 90),
        _pdf_line("EXPERIENCE", 120, bold=True),
        _pdf_line(
            "Acme Corp - Software Engineer    Jan 2020 – Present",
            140,
            [
                _run("Acme Corp", bold=True),
                _run(" - Software Engineer    "),
                _run("Jan 2020 – Present", color="666666", x0=450),
            ],
        ),
        _pdf_line("• Built the thing", 156, x0=90),
        _pdf_line("EDUCATION", 180, bold=True),
        _pdf_line(
            "BS, State University    2016 – 2020",
            200,
            [_run("BS", bold=True), _run(", State University    ", italic=True), _run("2016 – 2020", x0=470)],
        ),
        _pdf_line("PROJECTS", 224, bold=True),
        _pdf_line(
            "Widgetizer | github.com/jane/widgetizer    2021",
            244,
            [_run("Widgetizer", bold=True), _run(" | github.com/jane/widgetizer    2021")],
        ),
        _pdf_line("A widget tool", 260),
    ]
    return PdfDoc(pages=[PdfPage(number=0, width=612.0, height=792.0, lines=lines)])


class FormatProfileCompat(unittest.TestCase):
    def test_profile_from_before_these_fields_loads_with_old_behavior(self):
        old = asdict(default_format_profile())
        for key in [k for k in old if k.startswith(("education_title", "project_title", "contact_"))]:
            if key != "contact_separator":
                del old[key]
        del old["entry_title_tail_first"]
        fmt = FormatProfile.from_dict(json.loads(json.dumps(old)))
        self.assertEqual(fmt.contact_layout, CONTACT_INLINE)
        self.assertIsNone(fmt.education_title_separator)
        self.assertIsNone(fmt.project_title_tail_sub)
        # Education/projects fall back to the experience separator and the
        # title-first order, as before.
        paras = _paras(render_docx(_content(), fmt))
        self.assertTrue(_para(paras, "State University").text.startswith("BS — State University"))
        self.assertTrue(_para(paras, "Widgetizer").text.startswith("Widgetizer — github.com"))
        self.assertTrue(_para(paras, "jane@x.com").text.startswith("jane@x.com | 555-1234"))

    def test_unknown_keys_from_a_newer_capturer_are_ignored(self):
        """The API loads profiles the dispatcher wrote; a newer dispatcher
        must not make an older API fall back to the default look."""
        data = json.loads(default_format_profile().to_json())
        data["entry_title_future_knob"] = True
        data["roles"]["name"]["kerning"] = 0.2
        data["section_header_roles"] = {"skills": {"bold": True, "glow": 1}}
        fmt = FormatProfile.from_dict(data)
        self.assertTrue(fmt.roles["name"].bold)
        self.assertTrue(fmt.section_header_roles["skills"].bold)


class TitleSeparator(unittest.TestCase):
    def test_variants(self):
        cases = [
            ("Engineer — Acme", "Engineer", "Acme", (" — ", False)),
            ("Acme - Engineer    2020", "Engineer", "Acme", (" - ", True)),
            ("Engineer, Acme", "Engineer", "Acme", (", ", False)),
            ("Engineer at Acme", "Engineer", "Acme", (" at ", False)),
            ("ACME CORP | Engineer", "Engineer", "Acme Corp", (" | ", True)),
            # Bracketed institution: "( … )" — the renderer adds the ")".
            ("In Progress: Certificate (Coursera) | Expected", "Certificate", "Coursera", (" (", False)),
            # A tech stack between name and date: the joining token.
            ("VATMissions | Go, React | 2026", "VATMissions", "2026", (" | ", False)),
        ]
        for text, head, tail, want in cases:
            with self.subTest(text=text):
                self.assertEqual(_title_separator(text, head, tail), want)

    def test_date_range_dash_is_not_a_title_separator(self):
        # Company on the next line: the only dash is inside the dates.
        self.assertIsNone(
            _title_separator("Software Engineer    Jan 2020 – Present", "Software Engineer", "Acme")
        )

    def test_strict_needs_both_parts(self):
        self.assertIsNone(_title_separator("Widgetizer | Go", "Widgetizer", "2026", strict=True))
        self.assertEqual(_title_separator("Widgetizer | Go", "Widgetizer", "2026"), (" | ", False))

    def test_empty_skills_category_has_no_separator(self):
        self.assertIsNone(_skills_category_split("• Agile & Scrum", "", "Agile & Scrum"))


class SectionScopedOrder(unittest.TestCase):
    """Finding B: the company-first order learned from experience must not
    flip education or project lines."""

    def _check(self, fmt):
        self.assertEqual(fmt.entry_title_separator, " - ")
        self.assertTrue(fmt.entry_title_tail_first)
        self.assertEqual(fmt.education_title_separator, ", ")
        self.assertFalse(fmt.education_title_tail_first)
        self.assertEqual(fmt.project_title_separator, " | ")
        self.assertFalse(fmt.project_title_tail_first)
        paras = _paras(render_docx(_content(), fmt))
        self.assertTrue(_para(paras, "Software Engineer").text.startswith("Acme Corp - Software Engineer"))
        self.assertTrue(_para(paras, "State University").text.startswith("BS, State University"))
        self.assertTrue(_para(paras, "Widgetizer").text.startswith("Widgetizer | github.com/jane/widgetizer"))

    def test_docx(self):
        self._check(capture_format_profile(_company_first_doc(), _profile()))

    def test_pdf(self):
        self._check(capture_format_profile_pdf(_company_first_pdf(), _profile()))

    def test_education_order_detected_on_its_own(self):
        doc = Document()
        doc.add_paragraph("Jane Doe")
        doc.add_paragraph("EDUCATION")
        doc.add_paragraph("State University - BS    2016 – 2020")
        profile = _profile()
        profile.experience = []
        profile.projects = []
        fmt = capture_format_profile(doc, profile)
        self.assertTrue(fmt.education_title_tail_first)
        paras = _paras(render_docx(_content(), fmt))
        self.assertTrue(_para(paras, "BS").text.startswith("State University - BS"))
        # Experience keeps the default title-first order.
        self.assertTrue(_para(paras, "Acme").text.startswith("Software Engineer"))

    def test_bracketed_institution_renders_closed(self):
        fmt = default_format_profile()
        fmt.education_title_separator = " ("
        paras = _paras(render_docx(_content(), fmt))
        self.assertTrue(_para(paras, "State University").text.startswith("BS (State University)"))

    def test_experience_not_found_does_not_crash_capture(self):
        """The profile's experience isn't in the document but a project
        line is — capture used to read an unset exp_text here."""
        doc = Document()
        doc.add_paragraph("Jane Doe")
        doc.add_paragraph("Widgetizer | github.com/jane/widgetizer")
        profile = _profile()
        profile.experience[0].position = "Not In Doc"
        profile.experience[0].company = "Nowhere Inc"
        fmt = capture_format_profile(doc, profile)
        self.assertEqual(fmt.entry_title_separator, " | ")


class HeadTailStyles(unittest.TestCase):
    """Finding C: styles follow the source's positions, so a bold company
    in a company-first source stays bold and the title stays plain."""

    def _check_experience(self, fmt):
        self.assertTrue(fmt.entry_title_tail_sub)
        self.assertTrue(fmt.roles[ROLE_ENTRY_TITLE].bold)
        self.assertFalse(fmt.roles[ROLE_ENTRY_SUB].bold)
        para = _para(_paras(render_docx(_content(), fmt)), "Software Engineer")
        self.assertEqual(para.runs[0].text, "Acme Corp")
        self.assertTrue(para.runs[0].font.bold)
        self.assertEqual(para.runs[1].text, " - Software Engineer")
        self.assertFalse(para.runs[1].font.bold)

    def test_docx_company_first(self):
        self._check_experience(capture_format_profile(_company_first_doc(), _profile()))

    def test_pdf_company_first_from_style_runs(self):
        fmt = capture_format_profile_pdf(_company_first_pdf(), _profile())
        self._check_experience(fmt)
        # The gray dates run is its own role.
        self.assertEqual(fmt.roles[ROLE_DATES].color, "666666")
        para = _para(_paras(render_docx(_content(), fmt)), "Software Engineer")
        self.assertEqual(str(para.runs[-1].font.color.rgb), "666666")

    def test_education_keeps_its_own_tail_look(self):
        # Education "BS" (bold) + ", State University" (italic): its own
        # roles, not the experience ones.
        for fmt in (
            capture_format_profile(_company_first_doc(), _profile()),
            capture_format_profile_pdf(_company_first_pdf(), _profile()),
        ):
            self.assertTrue(fmt.education_title_tail_sub)
            self.assertTrue(fmt.roles[ROLE_EDUCATION_SUB].italic)
            para = _para(_paras(render_docx(_content(), fmt)), "State University")
            self.assertTrue(para.runs[0].font.bold)
            self.assertTrue(para.runs[1].font.italic)
            self.assertFalse(para.runs[1].font.bold)

    def test_project_tail_follows_project_look_not_experience(self):
        fmt = default_format_profile()
        fmt.entry_title_tail_sub = True
        fmt.project_title_tail_sub = False
        para = _para(_paras(render_docx(_content(), fmt)), "Widgetizer")
        self.assertTrue(para.runs[1].font.bold)


class Dates(unittest.TestCase):
    """Finding F: a missing start date never leaves a dangling dash."""

    def test_dates(self):
        fmt = default_format_profile()
        self.assertEqual(_dates("", "December 2025", fmt), "December 2025")
        self.assertEqual(_dates(None, "December 2025", fmt), "December 2025")
        self.assertEqual(_dates("Jan 2020", None, fmt), "Jan 2020 – Present")
        self.assertEqual(_dates("Jan 2020", "Mar 2022", fmt), "Jan 2020 – Mar 2022")
        self.assertEqual(_dates("2026", "2026", fmt), "2026")
        self.assertEqual(_dates("", None, fmt), "")
        self.assertEqual(_dates("", "Present", fmt), "Present")
        fmt.dates_open_label = None
        self.assertEqual(_dates("Jan 2020", None, fmt), "Jan 2020")

    def test_render_education_without_start(self):
        content = _content()
        content.education[0].start_date = ""
        content.education[0].end_date = "December 2025"
        para = _para(_paras(render_docx(content, default_format_profile())), "State University")
        self.assertTrue(para.text.endswith("\tDecember 2025"))
        self.assertNotIn("– December", para.text)

    def test_render_entry_without_any_dates(self):
        content = _content()
        content.education[0].start_date = ""
        content.education[0].end_date = None
        para = _para(_paras(render_docx(content, default_format_profile())), "State University")
        self.assertEqual(para.text, "BS — State University")


class StackedContact(unittest.TestCase):
    """Finding G: a labelled, one-field-per-line contact block (#49's
    "Address: / Phone: / Email:") round-trips instead of collapsing into
    "email | phone"."""

    def test_docx_labelled_block_with_tab(self):
        doc = Document()
        doc.add_paragraph("Jane Doe")
        for label, value in (("Address:", "Denver, CO"), ("Phone:", "555-1234"), ("Email:", "jane@x.com")):
            para = doc.add_paragraph()
            para.paragraph_format.tab_stops.add_tab_stop(Inches(1.0), WD_TAB_ALIGNMENT.LEFT)
            _runs(para, (f"{label}\t", True, False), (value, False, False))
        doc.add_paragraph("EXPERIENCE")
        fmt = capture_format_profile(doc, _profile())
        self.assertEqual(fmt.contact_layout, CONTACT_STACKED)
        self.assertEqual(fmt.contact_order, ["location", "phone", "email"])
        self.assertEqual(
            fmt.contact_labels, {"location": "Address:", "phone": "Phone:", "email": "Email:"}
        )
        self.assertTrue(fmt.contact_label_bold)
        self.assertEqual(fmt.contact_value_tab_in, 1.0)
        paras = _paras(render_docx(_content(), fmt))
        self.assertEqual(
            [p.text for p in paras[1:4]],
            ["Address:\tDenver, CO", "Phone:\t555-1234", "Email:\tjane@x.com"],
        )
        self.assertTrue(paras[1].runs[0].font.bold)
        self.assertFalse(paras[1].runs[1].font.bold)
        self.assertAlmostEqual(paras[1].paragraph_format.tab_stops[0].position.inches, 1.0)

    def test_pdf_label_column_gives_tab(self):
        """Values starting their own run at one shared x (#49's layout)."""
        lines = [_pdf_line("Jane Doe", 40, size=20, bold=True, x0=165)]
        for i, (label, value) in enumerate(
            (("Address:", "Denver, CO"), ("Phone:", "555-1234"), ("Email:", "jane@x.com"))
        ):
            lines.append(
                _pdf_line(
                    f"{label}    {value}",
                    70 + i * 14,
                    [_run(f"{label}    ", bold=True, x0=165), _run(value, x0=228.4)],
                    x0=165,
                )
            )
        doc = PdfDoc(pages=[PdfPage(number=0, width=612.0, height=792.0, lines=lines)])
        fmt = capture_format_profile_pdf(doc, _profile())
        self.assertEqual(fmt.contact_layout, CONTACT_STACKED)
        self.assertTrue(fmt.contact_label_bold)
        self.assertEqual(fmt.contact_value_tab_in, 0.88)

    def test_unlabelled_stack_and_city_repeated_on_entry_lines(self):
        """One field per line, no labels; the city also appears on an
        entry line read earlier (a right sidebar read after the main
        column) — the copy beside email/phone wins."""
        lines = [
            _pdf_line("Jane Doe", 40, size=20, bold=True),
            _pdf_line("Acme Corp    Denver, CO", 70),
            _pdf_line("jane@x.com", 100, x0=440),
            _pdf_line("555-1234", 114, x0=440),
            _pdf_line("Denver, CO", 128, x0=440),
        ]
        doc = PdfDoc(pages=[PdfPage(number=0, width=612.0, height=792.0, lines=lines)])
        fmt = capture_format_profile_pdf(doc, _profile())
        self.assertEqual(fmt.contact_layout, CONTACT_STACKED)
        self.assertEqual(fmt.contact_order, ["email", "phone", "location"])
        self.assertEqual(fmt.contact_labels, {})
        paras = _paras(render_docx(_content(), fmt))
        self.assertEqual([p.text for p in paras[1:4]], ["jane@x.com", "555-1234", "Denver, CO"])

    def test_inline_contact_line_stays_inline(self):
        fmt = capture_format_profile(_company_first_doc(), _profile())
        self.assertEqual(fmt.contact_layout, CONTACT_INLINE)
        paras = _paras(render_docx(_content(), fmt))
        self.assertEqual(paras[1].text, "jane@x.com | 555-1234 | Denver, CO")


def _located_profile(*locations: str) -> _Profile:
    """Two Acme/Globex roles; the given locations go on them in order."""
    profile = _profile()
    profile.experience = [
        _Exp("Acme Corp", "Senior Engineer", "Jan 2020", None, [_Bullet("Built the thing")]),
        _Exp("Globex", "Engineer", "Jun 2017", "Dec 2019", [_Bullet("Shipped the other thing")]),
    ]
    for exp, location in zip(profile.experience, locations):
        exp.location = location
    return profile


def _located_content(*locations: str | None) -> ResumeContent:
    content = _content()
    content.experiences = [
        ResumeExperience("Acme Corp", "Senior Engineer", "Jan 2020", None, ["Built the thing"]),
        ResumeExperience("Globex", "Engineer", "Jun 2017", "Dec 2019", ["Shipped the other thing"]),
    ]
    for exp, location in zip(content.experiences, locations):
        exp.location = location
    return content


def _entry_doc(*entries) -> Document:
    """An experience section whose entry lines are each given as run
    specs (text, bold, italic); a str entry is a plain paragraph."""
    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("jane@x.com | 555-1234 | Denver, CO")
    doc.add_paragraph("EXPERIENCE")
    for entry in entries:
        if isinstance(entry, str):
            doc.add_paragraph(entry)
        else:
            _runs(doc.add_paragraph(), *entry)
    return doc


_PIPE_CHAIN = (
    (
        ("Senior Engineer", True, True),
        (" | ", False, True),
        ("Acme Corp | Remote                         Jan 2020 - Present", False, True),
    ),
    "• Built the thing",
    (
        ("Engineer", True, True),
        (" | ", False, True),
        ("Globex | Denver, CO                         Jun 2017 - Dec 2019", False, True),
    ),
    "• Shipped the other thing",
)


class ExperienceLocation(unittest.TestCase):
    """An experience's location ("Remote") renders where the source puts
    it: chained onto the entry line with the source's separator, or on its
    own line under the entry."""

    def test_entry_without_location_is_unchanged(self):
        para = _para(_paras(render_docx(_located_content(None, ""), default_format_profile())), "Acme")
        self.assertEqual(para.text, "Senior Engineer — Acme Corp\tJan 2020 – Present")

    def test_default_chains_location_with_the_title_separator(self):
        paras = _paras(render_docx(_located_content("Remote", None), default_format_profile()))
        self.assertEqual(
            _para(paras, "Acme").text, "Senior Engineer — Acme Corp — Remote\tJan 2020 – Present"
        )
        self.assertEqual(_para(paras, "Globex").text, "Engineer — Globex\tJun 2017 – Dec 2019")

    def test_default_avoids_word_and_bracket_separators(self):
        fmt = default_format_profile()
        for title_sep, expected in (
            (" at ", "Senior Engineer at Acme Corp, Remote"),
            (" (", "Senior Engineer (Acme Corp), Remote"),
        ):
            fmt.entry_title_separator = title_sep
            para = _para(_paras(render_docx(_located_content("Remote"), fmt)), "Acme")
            self.assertEqual(para.text.split("\t")[0], expected)

    def test_location_follows_the_trailing_part_look(self):
        fmt = default_format_profile()
        fmt.entry_title_tail_sub = True
        para = _para(_paras(render_docx(_located_content("Remote"), fmt)), "Acme")
        location_run = next(r for r in para.runs if "Remote" in r.text)
        self.assertTrue(location_run.font.italic)
        self.assertFalse(location_run.font.bold)

    def test_profile_from_before_location_loads(self):
        old = asdict(default_format_profile())
        del old["entry_location_separator"]
        del old["entry_location_on_own_line"]
        old["entry_title_separator"] = " | "
        fmt = FormatProfile.from_dict(json.loads(json.dumps(old)))
        para = _para(_paras(render_docx(_located_content("Remote"), fmt)), "Acme")
        self.assertEqual(para.text, "Senior Engineer | Acme Corp | Remote\tJan 2020 – Present")

    def test_location_is_part_of_the_source_hash(self):
        fmt = default_format_profile()
        self.assertNotEqual(
            content_source_hash(_located_content(None), fmt=fmt),
            content_source_hash(_located_content("Remote"), fmt=fmt),
        )

    def test_docx_pipe_chain(self):
        """"Title | Company | Remote      dates": the location joins the
        chain with the source's own separator and look."""
        fmt = capture_format_profile(_entry_doc(*_PIPE_CHAIN), _located_profile("Remote", "Denver, CO"))
        self.assertEqual(fmt.entry_location_separator, " | ")
        self.assertFalse(fmt.entry_location_on_own_line)
        self.assertTrue(fmt.roles[ROLE_ENTRY_LOCATION].italic)
        self.assertFalse(fmt.roles[ROLE_ENTRY_LOCATION].bold)
        paras = _paras(render_docx(_located_content("Remote", "Denver, CO"), fmt))
        acme = _para(paras, "Acme")
        self.assertEqual(acme.text, "Senior Engineer | Acme Corp | Remote\tJan 2020 - Present")
        self.assertTrue(next(r for r in acme.runs if "Remote" in r.text).font.italic)
        self.assertEqual(
            _para(paras, "Globex").text, "Engineer | Globex | Denver, CO\tJun 2017 - Dec 2019"
        )

    def test_docx_comma_after_company(self):
        doc = _entry_doc(
            (("Senior Engineer", True, False), (" — Acme Corp, Denver, CO\tJan 2020 – Present", False, False)),
        )
        fmt = capture_format_profile(doc, _located_profile("Denver, CO"))
        self.assertEqual(fmt.entry_location_separator, ", ")
        para = _para(_paras(render_docx(_located_content("Denver, CO"), fmt)), "Acme")
        self.assertEqual(para.text, "Senior Engineer — Acme Corp, Denver, CO\tJan 2020 – Present")

    def test_docx_bracketed_location_closes(self):
        doc = _entry_doc(
            (("Senior Engineer at Acme Corp (Remote)\tJan 2020 – Present", False, False),),
        )
        fmt = capture_format_profile(doc, _located_profile("Remote"))
        self.assertEqual(fmt.entry_location_separator, " (")
        para = _para(_paras(render_docx(_located_content("Remote"), fmt)), "Acme")
        self.assertEqual(para.text, "Senior Engineer at Acme Corp (Remote)\tJan 2020 – Present")

    def test_docx_own_line_under_the_entry(self):
        doc = _entry_doc(
            (("Acme Corp", True, False), (" - Senior Engineer\tJan 2020 – Present", False, False)),
            (("Remote", False, True),),
            "• Built the thing",
        )
        fmt = capture_format_profile(doc, _located_profile("Remote"))
        self.assertTrue(fmt.entry_location_on_own_line)
        self.assertTrue(fmt.roles[ROLE_ENTRY_LOCATION].italic)
        paras = _paras(render_docx(_located_content("Remote"), fmt))
        entry = _para(paras, "Acme")
        self.assertEqual(entry.text, "Acme Corp - Senior Engineer\tJan 2020 – Present")
        below = paras[paras.index(entry) + 1]
        self.assertEqual(below.text, "Remote")
        self.assertTrue(below.runs[0].font.italic)

    def test_docx_own_line_comes_before_own_line_dates(self):
        fmt = default_format_profile()
        fmt.dates_on_title_line = False
        fmt.entry_location_on_own_line = True
        paras = _paras(render_docx(_located_content("Remote"), fmt))
        i = paras.index(_para(paras, "Acme"))
        self.assertEqual([p.text for p in paras[i : i + 3]], [
            "Senior Engineer — Acme Corp", "Remote", "Jan 2020 – Present",
        ])

    def test_location_only_on_a_later_role_is_still_found(self):
        doc = _entry_doc(
            (("Senior Engineer | Acme Corp\tJan 2020 – Present", False, False),),
            (("Engineer | Globex | Remote\tJun 2017 – Dec 2019", False, False),),
        )
        fmt = capture_format_profile(doc, _located_profile("", "Remote"))
        self.assertEqual(fmt.entry_location_separator, " | ")

    def test_entry_matched_only_after_whitespace_folding_does_not_crash(self):
        doc = _entry_doc(
            (("Senior  Engineer | Acme Corp | Remote\tJan 2020 – Present", False, False),),
        )
        profile = _located_profile("Remote")
        profile.experience[0].company = "Acme  Corp"
        fmt = capture_format_profile(doc, profile)
        self.assertIsNone(fmt.entry_location_separator)

    def test_location_on_the_right_gap_keeps_the_fallback(self):
        """A location floated right ("Acme Corp      Remote") isn't joined
        by a separator: the renderer's chained fallback stands in."""
        doc = _entry_doc(
            (("Senior Engineer — Acme Corp\tRemote", False, False),),
        )
        fmt = capture_format_profile(doc, _located_profile("Remote"))
        self.assertIsNone(fmt.entry_location_separator)
        self.assertFalse(fmt.entry_location_on_own_line)

    def test_pdf_pipe_chain(self):
        lines = [
            _pdf_line("Jane Doe", 60, size=22, bold=True),
            _pdf_line("jane@x.com | 555-1234 | Denver, CO", 90),
            _pdf_line("EXPERIENCE", 120, bold=True),
            _pdf_line(
                "Senior Engineer | Acme Corp | Remote    Jan 2020 – Present",
                140,
                [
                    _run("Senior Engineer", bold=True),
                    _run(" | Acme Corp | "),
                    _run("Remote", italic=True),
                    _run("    Jan 2020 – Present", x0=450),
                ],
            ),
            _pdf_line("• Built the thing", 156, x0=90),
        ]
        doc = PdfDoc(pages=[PdfPage(number=0, width=612.0, height=792.0, lines=lines)])
        fmt = capture_format_profile_pdf(doc, _located_profile("Remote"))
        self.assertEqual(fmt.entry_location_separator, " | ")
        self.assertFalse(fmt.entry_location_on_own_line)
        self.assertTrue(fmt.roles[ROLE_ENTRY_LOCATION].italic)
        para = _para(_paras(render_docx(_located_content("Remote"), fmt)), "Acme")
        self.assertEqual(para.text, "Senior Engineer | Acme Corp | Remote\tJan 2020 – Present")

    def test_pdf_own_line(self):
        lines = [
            _pdf_line("Jane Doe", 60, size=22, bold=True),
            _pdf_line("jane@x.com | 555-1234 | Denver, CO", 90),
            _pdf_line("EXPERIENCE", 120, bold=True),
            _pdf_line("Senior Engineer, Acme Corp    Jan 2020 – Present", 140, bold=True),
            _pdf_line("Remote", 156, [_run("Remote", italic=True)]),
            _pdf_line("• Built the thing", 172, x0=90),
        ]
        doc = PdfDoc(pages=[PdfPage(number=0, width=612.0, height=792.0, lines=lines)])
        fmt = capture_format_profile_pdf(doc, _located_profile("Remote"))
        self.assertTrue(fmt.entry_location_on_own_line)
        paras = _paras(render_docx(_located_content("Remote"), fmt))
        i = paras.index(_para(paras, "Acme"))
        self.assertEqual(paras[i + 1].text, "Remote")


if __name__ == "__main__":
    unittest.main()
