"""Fidelity tests against real resume fixtures.

fixtures/marquelle_original.docx is a real uploaded resume;
fixtures/marquelle_generated_bad.docx is what the pre-v2 pipeline emitted
from it (Calibri fallbacks, "—" separators, literal "● " glyphs, centered
italic headers). These tests pin capture + render to the original's layout
and assert the old defects stay gone.
"""

import io
import unittest
from dataclasses import dataclass, field
from pathlib import Path

import docx
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from shared.docgen import (
    ROLE_BODY,
    ROLE_CONTACT,
    ROLE_ENTRY_SUB,
    ROLE_ENTRY_TITLE,
    ROLE_LOCATION,
    ROLE_NAME,
    ROLE_SECTION_HEADER,
    SECTION_PROJECTS,
    SECTION_SKILLS,
    FormatProfile,
    ResumeContent,
    ResumeExperience,
    ResumeHeader,
    ResumeProject,
    ResumeSkillGroup,
    capture_format_profile,
    render_docx,
)

FIXTURES = Path(__file__).parent / "fixtures"
ORIGINAL = FIXTURES / "marquelle_original.docx"
GENERATED_BAD = FIXTURES / "marquelle_generated_bad.docx"


# Duck-typed stand-in for the parser's UserDocxProfile — capture only
# reads attributes.


@dataclass
class _Bullet:
    text: str


@dataclass
class _Basic:
    name: str = ""
    email: str = ""
    phone: str = ""
    location: str = ""
    links: list = field(default_factory=list)


@dataclass
class _Exp:
    company: str
    position: str
    start_date: str
    end_date: str | None
    bullets: list = field(default_factory=list)


@dataclass
class _Cat:
    category: str
    skills: list


@dataclass
class _Proj:
    name: str
    link: str | None
    description: str
    start_date: str
    end_date: str | None


@dataclass
class _Profile:
    summary: str | None
    basic_info: _Basic
    experience: list = field(default_factory=list)
    skills: list = field(default_factory=list)
    education: list = field(default_factory=list)
    projects: list = field(default_factory=list)


def _marquelle_profile() -> _Profile:
    """What the parse agent should extract from the original fixture."""
    return _Profile(
        summary=None,
        basic_info=_Basic(
            name="Marquelle Nesbitt",
            email="marquelle.nesbitt@gmail.com",
            location="Colorado, USA",
            links=["github.com/quellen-sol", "linkedin.com/in/quellen"],
        ),
        experience=[
            _Exp(
                company="Step Finance",
                position="Senior Backend Engineer",
                start_date="Apr 2022",
                end_date="Feb 2026",
                bullets=[
                    _Bullet(
                        "Led a team of 6 building a real-time blockchain "
                        "transaction ingestion platform with TypeScript, Rust, "
                        "PostgreSQL and Amazon SQS"
                    ),
                    _Bullet(
                        "Maintained a Node.js/TypeScript REST API for the "
                        "portfolio dashboard"
                    ),
                ],
            ),
            _Exp(
                company="SolDen",
                position="Full Stack Engineer (Contract)",
                start_date="Feb 2022",
                end_date="Apr 2022",
                bullets=[
                    _Bullet(
                        "Architected and maintained the backend using "
                        "TypeScript, Express, MongoDB and Redis"
                    )
                ],
            ),
        ],
        skills=[
            _Cat("Languages", ["TypeScript", "JavaScript", "Rust"]),
            _Cat("Infrastructure", ["AWS", "Terraform", "Docker"]),
        ],
        projects=[
            _Proj(
                name="VATMissions",
                link="github.com/quellen-sol/vatmissions",
                description="Career-style mission add-on for the VATSIM flight network",
                start_date="2026",
                end_date=None,
            )
        ],
    )


def _marquelle_content() -> ResumeContent:
    """The same resume as canonical ResumeContent for rendering."""
    return ResumeContent(
        header=ResumeHeader(
            name="Marquelle Nesbitt",
            email="marquelle.nesbitt@gmail.com",
            location="Colorado, USA",
            links=["github.com/quellen-sol", "linkedin.com/in/quellen"],
        ),
        experiences=[
            ResumeExperience(
                company="Step Finance",
                title="Senior Backend Engineer",
                start_date="Apr 2022",
                end_date="Feb 2026",
                bullets=[
                    (
                        "Led a team of 6 building a real-time blockchain "
                        "transaction ingestion platform with TypeScript, "
                        "Rust, PostgreSQL and Amazon SQS"
                    ),
                    (
                        "Maintained a Node.js/TypeScript REST API for the "
                        "portfolio dashboard"
                    ),
                ],
            ),
            ResumeExperience(
                company="SolDen",
                title="Full Stack Engineer (Contract)",
                start_date="Feb 2022",
                end_date="Apr 2022",
                bullets=[
                    (
                        "Architected and maintained the backend using "
                        "TypeScript, Express, MongoDB and Redis"
                    )
                ],
            ),
        ],
        skill_groups=[
            ResumeSkillGroup(
                category="Languages", skills=["TypeScript", "JavaScript", "Rust"]
            ),
            ResumeSkillGroup(
                category="Infrastructure", skills=["AWS", "Terraform", "Docker"]
            ),
        ],
        projects=[
            ResumeProject(
                name="VATMissions",
                link="github.com/quellen-sol/vatmissions",
                description="Career-style mission add-on for the VATSIM flight network",
                start_date="2026",
                end_date=None,
            )
        ],
    )


def _open(blob: bytes) -> Document:
    return docx.Document(io.BytesIO(blob))


def _size(run) -> float | None:
    return round(run.font.size.pt, 1) if run.font.size else None


class FixtureCapture(unittest.TestCase):
    """capture_format_profile on the real uploaded resume."""

    @classmethod
    def setUpClass(cls):
        cls.fmt = capture_format_profile(
            Document(ORIGINAL), _marquelle_profile()
        )

    def test_margins(self):
        m = self.fmt.page_margins_in
        for side in ("top", "right", "bottom", "left"):
            self.assertAlmostEqual(m[side], 0.625, places=2)

    def test_header_roles(self):
        name = self.fmt.roles[ROLE_NAME]
        self.assertEqual(name.font_name, "Arial")
        self.assertEqual(name.size_pt, 14.0)
        self.assertTrue(name.bold)
        self.assertEqual(name.color, "1F3864")
        self.assertEqual(name.alignment, "center")
        contact = self.fmt.roles[ROLE_CONTACT]
        self.assertEqual(contact.color, "444444")
        self.assertEqual(contact.alignment, "center")
        location = self.fmt.roles[ROLE_LOCATION]
        self.assertTrue(location.italic)
        self.assertEqual(location.color, "595959")
        self.assertTrue(self.fmt.location_on_own_line)

    def test_section_header_style(self):
        header = self.fmt.roles[ROLE_SECTION_HEADER]
        self.assertTrue(header.bold)
        self.assertEqual(header.font_name, "Arial")
        self.assertEqual(header.size_pt, 10.5)
        # Left-aligned, not the old centered-italic default.
        self.assertNotEqual(header.alignment, "center")
        self.assertFalse(header.italic)

    def test_entry_title_layout(self):
        self.assertEqual(self.fmt.entry_title_separator, " | ")
        # Title bold-italic, company tail plain italic.
        self.assertTrue(self.fmt.entry_title_tail_sub)
        self.assertTrue(self.fmt.dates_on_title_line)
        self.assertEqual(self.fmt.dates_separator, " - ")
        title = self.fmt.roles[ROLE_ENTRY_TITLE]
        self.assertTrue(title.bold)
        self.assertTrue(title.italic)
        sub = self.fmt.roles[ROLE_ENTRY_SUB]
        self.assertFalse(sub.bold)
        self.assertTrue(sub.italic)

    def test_bullet_look(self):
        # Numbering resolves to the real glyph, with hanging indent —
        # not a literal "● " prefix.
        self.assertEqual(self.fmt.bullet_glyph, "●")
        self.assertAlmostEqual(self.fmt.bullet_left_indent_in, 0.5)
        self.assertAlmostEqual(self.fmt.bullet_hanging_indent_in, 0.25)

    def test_skills_layout(self):
        self.assertEqual(self.fmt.skills_layout, "grouped")
        self.assertEqual(self.fmt.skills_category_separator, ": ")
        self.assertTrue(self.fmt.skills_category_bold)

    def test_project_layout(self):
        self.assertTrue(self.fmt.project_link_on_own_line)
        self.assertTrue(self.fmt.project_dates_inline)

    def test_section_blank_lines(self):
        self.assertEqual(
            self.fmt.section_blank_lines.get(SECTION_SKILLS), 4
        )
        self.assertEqual(
            self.fmt.section_blank_lines.get(SECTION_PROJECTS), 1
        )

    def test_body_style_from_items_not_label(self):
        # Skills lines split into a bold label and plain items — the body
        # look comes from the items run (9.5pt, not bold).
        body = self.fmt.roles[ROLE_BODY]
        self.assertFalse(body.bold)
        self.assertEqual(body.size_pt, 9.5)

    def test_no_foreign_fonts(self):
        for role, style in self.fmt.roles.items():
            self.assertEqual(
                style.font_name, "Arial", f"role {role} should be Arial"
            )

    def test_profile_serializes(self):
        self.assertEqual(
            FormatProfile.from_json(self.fmt.to_json()), self.fmt
        )


class FixtureRender(unittest.TestCase):
    """render_docx through the captured profile reproduces the original's
    layout and sheds the defects visible in generated_bad.docx."""

    @classmethod
    def setUpClass(cls):
        cls.fmt = capture_format_profile(
            Document(ORIGINAL), _marquelle_profile()
        )
        cls.doc = _open(render_docx(_marquelle_content(), cls.fmt))
        cls.paras = cls.doc.paragraphs

    def _by_text(self, startswith: str):
        for p in self.paras:
            text = p.text.strip().lstrip("●\t•·-").lstrip()
            if text.startswith(startswith):
                return p
        self.fail(f"no paragraph starting with {startswith!r}")

    def test_name_and_contact_lines(self):
        name = self.paras[0]
        self.assertEqual(name.text, "Marquelle Nesbitt")
        self.assertTrue(name.runs[0].font.bold)
        self.assertEqual(name.runs[0].font.name, "Arial")
        self.assertEqual(_size(name.runs[0]), 14.0)
        contact = self.paras[1]
        self.assertIn(" | ", contact.text)
        self.assertIn("github.com/quellen-sol", contact.text)
        self.assertIn("linkedin.com/in/quellen", contact.text)
        self.assertEqual(str(contact.runs[0].font.color.rgb), "444444")

    def test_location_own_line(self):
        loc = self.paras[2]
        self.assertEqual(loc.text, "Colorado, USA")
        self.assertTrue(loc.runs[0].font.italic)
        self.assertEqual(str(loc.runs[0].font.color.rgb), "595959")

    def test_title_line_uses_pipe_and_hyphen(self):
        title = self._by_text("Senior Backend Engineer")
        self.assertIn(" | Step Finance", title.text)
        self.assertIn("Apr 2022 - Feb 2026", title.text)
        self.assertNotIn(" — ", title.text)
        self.assertNotIn(" – ", title.text)
        head, tail = title.runs[0], title.runs[-1]
        self.assertTrue(head.font.bold)
        self.assertTrue(head.font.italic)
        self.assertFalse(tail.font.bold)
        self.assertTrue(tail.font.italic)

    def test_bullets_are_hanging_indent(self):
        bullet = self._by_text("Led a team of 6")
        pf = bullet.paragraph_format
        self.assertAlmostEqual(pf.left_indent.inches, 0.5)
        self.assertAlmostEqual(pf.first_line_indent.inches, -0.25)
        # Glyph + tab, never "● " baked into the text.
        self.assertEqual(bullet.runs[0].text, "●\t")
        self.assertFalse(bullet.text.startswith("● "))

    def test_section_blank_lines_preserved(self):
        skills_idx = next(
            i for i, p in enumerate(self.paras) if p.text == "SKILLS"
        )
        above = self.paras[skills_idx - 4 : skills_idx]
        self.assertEqual(len(above), 4)
        for p in above:
            self.assertFalse(p.text.strip())
        projects_idx = next(
            i for i, p in enumerate(self.paras) if p.text == "PROJECTS"
        )
        self.assertFalse(self.paras[projects_idx - 1].text.strip())
        self.assertTrue(
            self.paras[projects_idx - 2].text.strip()
        )

    def test_skills_bold_labels(self):
        line = self._by_text("Languages")
        label, items = line.runs[0], line.runs[-1]
        self.assertTrue(label.font.bold)
        self.assertFalse(items.font.bold)
        self.assertEqual(label.text, "Languages: ")
        self.assertEqual(_size(label), 9.5)

    def test_project_layout(self):
        title = self._by_text("VATMissions")
        self.assertIn(" | ", title.text)
        link = self._by_text("github.com/quellen-sol/vatmissions")
        self.assertNotIn("VATMissions", link.text)

    def test_no_calibri_anywhere(self):
        for p in self.paras:
            for run in p.runs:
                self.assertNotEqual(
                    run.font.name, "Calibri", f"run {run.text!r}"
                )

    def test_margins_round_trip(self):
        section = self.doc.sections[0]
        self.assertAlmostEqual(section.top_margin.inches, 0.625, places=2)


class GeneratedBadRegression(unittest.TestCase):
    """Pin the specific defects of the old generated.docx so they stay
    gone: Calibri fallback, centered italic headers, literal "● " glyph
    runs, em/en dash separators, duplicated location."""

    @classmethod
    def setUpClass(cls):
        cls.fmt = capture_format_profile(
            Document(ORIGINAL), _marquelle_profile()
        )
        cls.doc = _open(render_docx(_marquelle_content(), cls.fmt))
        cls.paras = cls.doc.paragraphs
        cls.texts = [p.text for p in cls.doc.paragraphs if p.text.strip()]
        cls.bad = Document(GENERATED_BAD)

    def test_bad_fixture_shows_old_defects(self):
        # Guard: the fixture really is the old bad output, so these
        # assertions mean something.
        bad_join = "\n".join(p.text for p in self.bad.paragraphs)
        self.assertIn("● ", bad_join)
        self.assertIn(" — ", bad_join)

    def test_headers_not_centered_italic(self):
        for p in self.paras:
            if p.text.strip() in ("EXPERIENCE", "SKILLS", "PROJECTS"):
                self.assertNotEqual(
                    p.alignment, WD_ALIGN_PARAGRAPH.CENTER,
                    f"{p.text} should not be centered",
                )
                for run in p.runs:
                    self.assertFalse(run.font.italic)
                    self.assertTrue(run.font.bold)

    def test_no_literal_bullet_glyph_space(self):
        for p in self.paras:
            self.assertFalse(
                p.text.lstrip().startswith("● "), f"{p.text!r}"
            )

    def test_no_em_dash_separator(self):
        joined = "\n".join(self.texts)
        self.assertNotIn(" — ", joined)
        self.assertNotIn(" – ", joined)

    def test_location_not_duplicated(self):
        count = sum("Colorado" in t for t in self.texts)
        self.assertEqual(count, 1)


if __name__ == "__main__":
    unittest.main()
