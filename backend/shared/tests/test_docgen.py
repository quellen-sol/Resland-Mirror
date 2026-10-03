"""Docgen capture/render tests — stdlib unittest, no extra deps.

Run: .venv/bin/python -m unittest backend.shared.tests.test_docgen -v
     (from the repo root with the workspace venv)
"""

import io
import json
import unittest
from dataclasses import dataclass, field

import docx
from docx import Document
from docx.shared import Pt

from shared.docgen import (
    ALL_ROLES,
    ALL_SECTIONS,
    ROLE_BODY,
    ROLE_NAME,
    ROLE_SECTION_HEADER,
    SECTION_EDUCATION,
    SECTION_EXPERIENCE,
    SECTION_SKILLS,
    SECTION_SUMMARY,
    FormatProfile,
    ResumeContent,
    _section_header_idx,
    ResumeEducation,
    ResumeExperience,
    ResumeHeader,
    ResumeProject,
    ResumeSkillGroup,
    capture_format_profile,
    capture_format_profile_pdf,
    content_source_hash,
    default_format_profile,
    render_docx,
)
from shared.pdfdoc import PdfDoc, PdfLine, PdfPage, PdfStyle


# Duck-typed stand-in for the parser's UserDocxProfile — capture only reads
# attributes, so anything shaped the same works.


@dataclass
class _Basic:
    name: str = ""
    email: str = ""
    phone: str = ""
    location: str = ""


@dataclass
class _Bullet:
    text: str


@dataclass
class _Exp:
    company: str
    position: str
    start_date: str
    end_date: str | None
    bullets: list = field(default_factory=list)


@dataclass
class _Cat:
    category: str
    skills: list


@dataclass
class _Edu:
    institution: str
    degree: str
    field_of_study: str
    start_date: str
    end_date: str | None
    description: str = ""


@dataclass
class _Proj:
    name: str
    link: str | None
    description: str
    start_date: str
    end_date: str | None


@dataclass
class _Profile:
    summary: str | None
    basic_info: _Basic
    experience: list = field(default_factory=list)
    skills: list = field(default_factory=list)
    education: list = field(default_factory=list)
    projects: list = field(default_factory=list)


def _sample_profile() -> _Profile:
    return _Profile(
        summary="Engineer who ships.",
        basic_info=_Basic(
            name="Jane Doe", email="jane@x.com", phone="555-1", location="NYC"
        ),
        experience=[
            _Exp(
                company="Acme",
                position="Senior Engineer",
                start_date="Jan 2020",
                end_date=None,
                bullets=[_Bullet("Built the thing"), _Bullet("Shipped it too")],
            )
        ],
        skills=[_Cat(category="Languages", skills=["Python", "Go"])],
        education=[
            _Edu(
                institution="MIT",
                degree="BS",
                field_of_study="CS",
                start_date="2016",
                end_date="2020",
            )
        ],
        projects=[
            _Proj(
                name="Side",
                link="side.dev",
                description="toy",
                start_date="2021",
                end_date=None,
            )
        ],
    )


def _resume_doc() -> Document:
    doc = Document()
    name_para = doc.add_paragraph()
    run = name_para.add_run("Jane Doe")
    run.font.size = Pt(24)
    run.font.bold = True
    doc.add_paragraph("jane@x.com | 555-1 | NYC")
    doc.add_paragraph("SUMMARY")
    doc.add_paragraph("Engineer who ships.")
    doc.add_paragraph("EXPERIENCE")
    doc.add_paragraph("Senior Engineer — Acme    Jan 2020 – Present")
    doc.add_paragraph("• Built the thing")
    doc.add_paragraph("• Shipped it too")
    doc.add_paragraph("SKILLS")
    doc.add_paragraph("Languages: Python, Go")
    doc.add_paragraph("EDUCATION")
    doc.add_paragraph("BS — MIT    2016 – 2020")
    doc.add_paragraph("CS")
    return doc


def _sample_content() -> ResumeContent:
    return ResumeContent(
        header=ResumeHeader(
            name="Jane Doe",
            email="jane@x.com",
            phone="555-1",
            location="NYC",
            links=["github.com/jane"],
        ),
        summary="Engineer who ships.",
        experiences=[
            ResumeExperience(
                company="Acme",
                title="Senior Engineer",
                start_date="Jan 2020",
                end_date=None,
                bullets=["Built the thing", "Shipped it too"],
            )
        ],
        skill_groups=[ResumeSkillGroup(category="Languages", skills=["Python", "Go"])],
        education=[
            ResumeEducation(
                institution="MIT",
                degree="BS",
                field_of_study="CS",
                start_date="2016",
                end_date="2020",
            )
        ],
        projects=[],
    )


class FormatProfileSerialization(unittest.TestCase):
    def test_json_roundtrip(self):
        fmt = default_format_profile()
        restored = FormatProfile.from_json(fmt.to_json())
        self.assertEqual(restored, fmt)

    def test_from_dict_restores_role_styles(self):
        fmt = default_format_profile()
        restored = FormatProfile.from_dict(json.loads(fmt.to_json()))
        self.assertEqual(
            restored.roles[ROLE_NAME].size_pt, fmt.roles[ROLE_NAME].size_pt
        )


class Capture(unittest.TestCase):
    def setUp(self):
        self.fmt = capture_format_profile(_resume_doc(), _sample_profile())

    def test_captures_all_roles(self):
        for role in ALL_ROLES:
            self.assertIn(role, self.fmt.roles, f"missing role {role}")

    def test_captures_name_style(self):
        self.assertEqual(self.fmt.roles[ROLE_NAME].size_pt, 24.0)
        self.assertTrue(self.fmt.roles[ROLE_NAME].bold)

    def test_captures_section_order_and_titles(self):
        self.assertEqual(
            self.fmt.section_order[:4],
            [
                SECTION_SUMMARY,
                SECTION_EXPERIENCE,
                SECTION_SKILLS,
                SECTION_EDUCATION,
            ],
        )
        self.assertEqual(self.fmt.section_titles[SECTION_EXPERIENCE], "EXPERIENCE")

    def test_captures_bullet_glyph(self):
        self.assertEqual(self.fmt.bullet_glyph, "•")

    def test_captures_contact_separator(self):
        self.assertEqual(self.fmt.contact_separator, "|")

    def test_serializes_to_json(self):
        # The DB stores fmt.to_json() — it must round-trip without loss.
        self.assertEqual(FormatProfile.from_json(self.fmt.to_json()), self.fmt)


class SectionHeaderIdx(unittest.TestCase):
    class _P:
        """Duck-typed paragraph — the scan only reads .text."""

        def __init__(self, text):
            self.text = text

    def test_punctuated_tail_does_not_shadow_header(self):
        paras = [
            self._P("EXPERIENCE"),
            self._P(""),
            self._P("A stray short line."),
            self._P("First entry title"),
        ]
        self.assertEqual(_section_header_idx(paras, 3, set()), 0)

    def test_punctuated_only_gives_none(self):
        paras = [self._P("A stray short line."), self._P("First entry")]
        self.assertIsNone(_section_header_idx(paras, 1, set()))


class Render(unittest.TestCase):
    def _texts(self, blob: bytes) -> list[str]:
        doc = docx.Document(io.BytesIO(blob))
        return [p.text for p in doc.paragraphs if p.text.strip()]

    def test_roundtrip_contains_content(self):
        content = _sample_content()
        blob = render_docx(content, default_format_profile())
        texts = self._texts(blob)
        joined = "\n".join(texts)
        for expected in [
            "Jane Doe",
            "jane@x.com",
            "github.com/jane",
            "SUMMARY",
            "Engineer who ships.",
            "EXPERIENCE",
            "Senior Engineer — Acme",
            "Jan 2020 – Present",
            "• Built the thing",
            "SKILLS",
            "Python, Go",
            "EDUCATION",
            "BS — MIT",
        ]:
            self.assertIn(expected, joined)

    def test_section_order_follows_format(self):
        content = _sample_content()
        fmt = default_format_profile()
        fmt.section_order = [
            SECTION_EDUCATION,
            SECTION_EXPERIENCE,
            SECTION_SKILLS,
            SECTION_SUMMARY,
        ]
        texts = self._texts(render_docx(content, fmt))
        self.assertLess(texts.index("EDUCATION"), texts.index("EXPERIENCE"))

    def test_skips_empty_sections(self):
        content = _sample_content()
        content.summary = None
        content.projects = []
        texts = self._texts(render_docx(content, default_format_profile()))
        self.assertNotIn("SUMMARY", texts)
        self.assertNotIn("PROJECTS", texts)

    def test_grouped_skills_render_categories(self):
        content = _sample_content()
        content.skill_groups = [
            ResumeSkillGroup(category="Languages", skills=["Python"]),
            ResumeSkillGroup(category="Tools", skills=["Docker"]),
        ]
        fmt = default_format_profile()
        fmt.skills_layout = "grouped"
        texts = "\n".join(self._texts(render_docx(content, fmt)))
        self.assertIn("Languages: Python", texts)
        self.assertIn("Tools: Docker", texts)


class SourceHash(unittest.TestCase):
    def test_stable_for_identical_content(self):
        self.assertEqual(
            content_source_hash(_sample_content()),
            content_source_hash(_sample_content()),
        )

    def test_changes_when_format_profile_changes(self):
        # A re-upload that only changes layout must invalidate the export.
        fmt_a = default_format_profile()
        fmt_b = default_format_profile()
        fmt_b.section_blank_lines = {"skills": 4}
        self.assertNotEqual(
            content_source_hash(_sample_content(), fmt=fmt_a),
            content_source_hash(_sample_content(), fmt=fmt_b),
        )
        self.assertNotEqual(
            content_source_hash(_sample_content()),
            content_source_hash(_sample_content(), fmt=fmt_a),
        )

    def test_changes_when_content_changes(self):
        a = _sample_content()
        b = _sample_content()
        b.experiences[0].bullets[0] = "Rebuilt the thing"
        self.assertNotEqual(content_source_hash(a), content_source_hash(b))

    def test_changes_when_summary_dropped(self):
        a = _sample_content()
        b = _sample_content()
        b.summary = None
        self.assertNotEqual(content_source_hash(a), content_source_hash(b))

    def test_changes_with_template(self):
        self.assertNotEqual(
            content_source_hash(_sample_content(), template="original"),
            content_source_hash(_sample_content(), template="other"),
        )

    def test_changes_with_source_tag(self):
        a = _sample_content()
        b = _sample_content()
        b.source = "tailor_run:9"
        self.assertNotEqual(content_source_hash(a), content_source_hash(b))


class FormatProfileDefaults(unittest.TestCase):
    def test_default_profile_covers_all_roles(self):
        fmt = default_format_profile()
        for role in ALL_ROLES:
            self.assertIn(role, fmt.roles)

    def test_role_accessor_fills_missing(self):
        fmt = FormatProfile()
        # Unset roles come back as a blank style, not KeyError/None.
        self.assertIsNotNone(fmt.role(ROLE_NAME))
        self.assertFalse(fmt.role(ROLE_NAME).bold)


class CaptureEdgeCases(unittest.TestCase):
    def test_name_falls_back_to_first_nonempty_paragraph(self):
        # A profile name that never appears in the doc → the first
        # non-empty paragraph is treated as the name.
        doc = Document()
        name_para = doc.add_paragraph()
        run = name_para.add_run("Jane Doe")
        run.font.size = Pt(24)
        doc.add_paragraph("jane@x.com | 555-1 | NYC")

        profile = _sample_profile()
        profile.basic_info.name = "Nobody In Doc"
        fmt = capture_format_profile(doc, profile)
        self.assertEqual(fmt.roles[ROLE_NAME].size_pt, 24.0)

    def test_unmatched_sections_fall_back_to_default_titles(self):
        doc = Document()
        doc.add_paragraph("Jane Doe")
        doc.add_paragraph("jane@x.com")

        profile = _Profile(summary=None, basic_info=_Basic(name="Jane Doe"))
        fmt = capture_format_profile(doc, profile)
        self.assertEqual(fmt.section_titles[SECTION_EXPERIENCE], "EXPERIENCE")
        # Unmatched sections are appended, not dropped from the order.
        for section in ALL_SECTIONS:
            self.assertIn(section, fmt.section_order)

    def test_dates_on_own_line_detected(self):
        doc = Document()
        doc.add_paragraph("Jane Doe")
        doc.add_paragraph("jane@x.com")
        doc.add_paragraph("EXPERIENCE")
        doc.add_paragraph("Senior Engineer — Acme")
        doc.add_paragraph("Jan 2020 – Present")
        doc.add_paragraph("• Built the thing")

        fmt = capture_format_profile(doc, _sample_profile())
        self.assertFalse(fmt.dates_on_title_line)
        # The separate date line still tells capture its separator and the
        # open-ended label — exports must keep "– Present".
        self.assertEqual(fmt.dates_separator, " – ")
        self.assertEqual(fmt.dates_open_label, "Present")

    def test_year_only_dates_separator(self):
        doc = Document()
        doc.add_paragraph("Jane Doe")
        doc.add_paragraph("jane@x.com")
        doc.add_paragraph("EDUCATION")
        doc.add_paragraph("BS — MIT    2016 - 2020")

        profile = _Profile(
            summary=None,
            basic_info=_Basic(name="Jane Doe", email="jane@x.com"),
            education=[
                _Edu(
                    institution="MIT", degree="BS", field_of_study="CS",
                    start_date="2016", end_date="2020",
                )
            ],
        )
        fmt = capture_format_profile(doc, profile)
        self.assertEqual(fmt.dates_separator, " - ")

    def test_inline_project_dates_keep_inline_link(self):
        # "Name | link | dates" — both tail segments on the title line.
        doc = Document()
        doc.add_paragraph("Jane Doe")
        doc.add_paragraph("jane@x.com")
        doc.add_paragraph("PROJECTS")
        doc.add_paragraph("Demo | demo.dev | 2026")

        profile = _Profile(
            summary=None,
            basic_info=_Basic(name="Jane Doe", email="jane@x.com"),
            projects=[
                _Proj(
                    name="Demo", link="demo.dev", description="d",
                    start_date="2026", end_date=None,
                )
            ],
        )
        fmt = capture_format_profile(doc, profile)
        self.assertTrue(fmt.project_dates_inline)
        self.assertFalse(fmt.project_link_on_own_line)

        content = _sample_content()
        content.projects = [
            ResumeProject(
                name="Demo", link="demo.dev", description="d",
                start_date="2026", end_date=None,
            )
        ]
        texts = "\n".join(
            p.text
            for p in docx.Document(io.BytesIO(render_docx(content, fmt))).paragraphs
            if p.text.strip()
        )
        self.assertIn("Demo | demo.dev | 2026", texts)

    def test_numbered_list_bullet_detected(self):
        doc = Document()
        doc.add_paragraph("Jane Doe")
        doc.add_paragraph("jane@x.com")
        doc.add_paragraph("EXPERIENCE")
        doc.add_paragraph("Senior Engineer — Acme    Jan 2020 – Present")
        # No literal glyph, but a list style → treated as a bulleted line.
        doc.add_paragraph("Built the thing", style="List Bullet")

        fmt = capture_format_profile(doc, _sample_profile())
        self.assertEqual(fmt.bullet_glyph, "•")

    def test_contact_separator_variants(self):
        for sep in ("|", "•", "·", "-"):
            doc = Document()
            doc.add_paragraph("Jane Doe")
            doc.add_paragraph(f"jane@x.com {sep} 555-1 {sep} NYC")
            fmt = capture_format_profile(
                doc, _Profile(summary=None, basic_info=_Basic(
                    name="Jane Doe", email="jane@x.com", phone="555-1"))
            )
            self.assertEqual(fmt.contact_separator, sep, f"separator {sep!r}")

    def test_single_skill_category_is_flat_layout(self):
        fmt = capture_format_profile(_resume_doc(), _sample_profile())
        self.assertEqual(fmt.skills_layout, "flat")

    def test_multiple_skill_categories_are_grouped(self):
        profile = _sample_profile()
        profile.skills.append(_Cat(category="Tools", skills=["Docker"]))
        doc = _resume_doc()
        doc.add_paragraph("Tools: Docker")
        fmt = capture_format_profile(doc, profile)
        self.assertEqual(fmt.skills_layout, "grouped")

    def test_empty_document_produces_defaults(self):
        fmt = capture_format_profile(Document(), _sample_profile())
        for role in ALL_ROLES:
            self.assertIn(role, fmt.roles)
        self.assertEqual(fmt.columns, 1)


def _pdf_line(
    text,
    *,
    x0=72.0,
    x1=300.0,
    top=100.0,
    font="Helvetica",
    size=11.0,
    bold=False,
    color=None,
    gap=0.0,
):
    return PdfLine(
        text=text,
        page=0,
        x0=x0,
        x1=x1,
        top=top,
        bottom=top + size + 2,
        style=PdfStyle(font_name=font, size_pt=size, bold=bold, color=color),
        gap_before_pt=gap,
    )


def _resume_pdf_doc() -> PdfDoc:
    """Same resume as _resume_doc(), as pdf IR: Helvetica 11 body, bold 24
    name, bold 14 headers — lines carry geometry instead of docx styles."""
    top = 72.0

    def advance(lines, text, *, size=11.0, gap=0.0, **kw):
        nonlocal top
        top += gap
        lines.append(_pdf_line(text, top=top, size=size, gap=gap, **kw))
        top += size + 4

    lines = []
    advance(lines, "Jane Doe", font="Helvetica-Bold", size=24, bold=True, x1=200)
    advance(lines, "jane@x.com | 555-1 | NYC", x1=320)
    advance(lines, "SUMMARY", font="Helvetica-Bold", size=14, bold=True, gap=18)
    advance(lines, "Engineer who ships.", x1=400)
    advance(lines, "EXPERIENCE", font="Helvetica-Bold", size=14, bold=True, gap=18)
    advance(lines, "Senior Engineer — Acme    Jan 2020 – Present", x1=540)
    advance(lines, "• Built the thing", x0=90)
    advance(lines, "• Shipped it too", x0=90)
    advance(lines, "SKILLS", font="Helvetica-Bold", size=14, bold=True, gap=18)
    advance(lines, "Languages: Python, Go", x1=400)
    advance(lines, "EDUCATION", font="Helvetica-Bold", size=14, bold=True, gap=18)
    advance(lines, "BS — MIT    2016 – 2020", x1=500)
    advance(lines, "CS")
    return PdfDoc(
        pages=[PdfPage(number=0, width=612.0, height=792.0, lines=lines)]
    )


class PdfCapture(unittest.TestCase):
    def setUp(self):
        self.fmt = capture_format_profile_pdf(_resume_pdf_doc(), _sample_profile())

    def test_captures_all_roles(self):
        for role in ALL_ROLES:
            self.assertIn(role, self.fmt.roles, f"missing role {role}")

    def test_captures_name_style(self):
        self.assertEqual(self.fmt.roles[ROLE_NAME].size_pt, 24.0)
        self.assertTrue(self.fmt.roles[ROLE_NAME].bold)
        self.assertEqual(self.fmt.roles[ROLE_NAME].font_name, "Helvetica-Bold")

    def test_captures_header_style_and_gap(self):
        header = self.fmt.roles[ROLE_SECTION_HEADER]
        self.assertEqual(header.size_pt, 14.0)
        self.assertTrue(header.bold)
        self.assertEqual(header.space_before_pt, 18.0)

    def test_captures_body_style(self):
        self.assertEqual(self.fmt.roles[ROLE_BODY].font_name, "Helvetica")
        self.assertEqual(self.fmt.roles[ROLE_BODY].size_pt, 11.0)

    def test_captures_section_order_and_titles(self):
        self.assertEqual(
            self.fmt.section_order[:4],
            [
                SECTION_SUMMARY,
                SECTION_EXPERIENCE,
                SECTION_SKILLS,
                SECTION_EDUCATION,
            ],
        )
        self.assertEqual(self.fmt.section_titles[SECTION_EXPERIENCE], "EXPERIENCE")

    def test_captures_bullet_glyph_and_indent(self):
        self.assertEqual(self.fmt.bullet_glyph, "•")
        # Bullets sit 18pt right of the entry line → 0.25in.
        self.assertAlmostEqual(self.fmt.bullet_left_indent_in, 0.25)

    def test_captures_margins_from_geometry(self):
        m = self.fmt.page_margins_in
        self.assertEqual(m["left"], 1.0)
        self.assertEqual(m["top"], 1.0)
        # Content reaches x=540 on a 612pt page → 1.0in right margin.
        self.assertEqual(m["right"], 1.0)

    def test_captures_contact_separator(self):
        self.assertEqual(self.fmt.contact_separator, "|")

    def test_dates_on_title_line_detected(self):
        self.assertTrue(self.fmt.dates_on_title_line)

    def test_serializes_to_json(self):
        self.assertEqual(FormatProfile.from_json(self.fmt.to_json()), self.fmt)


class PdfCaptureEdgeCases(unittest.TestCase):
    def test_dates_on_own_line_detected(self):
        lines = [
            _pdf_line("Jane Doe", top=72),
            _pdf_line("jane@x.com", top=100),
            _pdf_line("EXPERIENCE", top=140, gap=18, bold=True),
            _pdf_line("Senior Engineer — Acme", top=170),
            _pdf_line("Jan 2020 – Present", top=188),
            _pdf_line("• Built the thing", top=206, x0=90),
        ]
        doc = PdfDoc(
            pages=[PdfPage(number=0, width=612.0, height=792.0, lines=lines)]
        )
        fmt = capture_format_profile_pdf(doc, _sample_profile())
        self.assertFalse(fmt.dates_on_title_line)

    def test_two_column_page_detected(self):
        lines = []
        for i in range(6):
            lines.append(_pdf_line(f"side {i}", x0=72, x1=200, top=100 + i * 20))
            lines.append(_pdf_line(f"main {i}", x0=230, x1=540, top=100 + i * 20))
        doc = PdfDoc(
            pages=[PdfPage(number=0, width=612.0, height=792.0, lines=lines)]
        )
        fmt = capture_format_profile_pdf(doc, _sample_profile())
        self.assertEqual(fmt.columns, 2)

    def test_empty_pdf_produces_defaults(self):
        fmt = capture_format_profile_pdf(PdfDoc(), _sample_profile())
        for role in ALL_ROLES:
            self.assertIn(role, fmt.roles)
        self.assertEqual(fmt.columns, 1)
        self.assertEqual(fmt.page_margins_in["left"], 1.0)

    def test_centered_name_alignment(self):
        # Name midpoint sits on the content column center; the wide body
        # line below establishes the column's edges.
        lines = [
            _pdf_line("Jane Doe", x0=256, x1=356, top=72, bold=True, size=24),
            _pdf_line("jane@x.com", x0=72, x1=200, top=110),
            _pdf_line("Engineer who ships.", x0=72, x1=540, top=140),
        ]
        doc = PdfDoc(
            pages=[PdfPage(number=0, width=612.0, height=792.0, lines=lines)]
        )
        fmt = capture_format_profile_pdf(
            doc, _Profile(summary=None, basic_info=_Basic(name="Jane Doe", email="jane@x.com"))
        )
        self.assertEqual(fmt.roles[ROLE_NAME].alignment, "center")


class RenderEdgeCases(unittest.TestCase):
    def _texts(self, blob: bytes) -> list[str]:
        doc = docx.Document(io.BytesIO(blob))
        return [p.text for p in doc.paragraphs if p.text.strip()]

    def test_dates_on_own_line_render_as_two_paragraphs(self):
        content = _sample_content()
        fmt = default_format_profile()
        fmt.dates_on_title_line = False
        texts = self._texts(render_docx(content, fmt))
        self.assertIn("Senior Engineer — Acme", texts)
        self.assertIn("Jan 2020 – Present", texts)

    def test_flat_skills_render_single_line(self):
        content = _sample_content()
        content.skill_groups = [
            ResumeSkillGroup(category=None, skills=["Python", "Go"])
        ]
        fmt = default_format_profile()
        fmt.skills_layout = "flat"
        texts = "\n".join(self._texts(render_docx(content, fmt)))
        self.assertIn("Python, Go", texts)

    def test_projects_render_with_link_and_description(self):
        content = _sample_content()
        content.projects = [
            ResumeProject(
                name="Side",
                link="side.dev",
                description="toy project",
                start_date="2021",
                end_date=None,
            )
        ]
        texts = "\n".join(
            self._texts(render_docx(content, default_format_profile()))
        )
        self.assertIn("PROJECTS", texts)
        self.assertIn("Side — side.dev", texts)
        self.assertIn("toy project", texts)

    def test_no_bullet_glyph_renders_plain_text(self):
        content = _sample_content()
        fmt = default_format_profile()
        fmt.bullet_glyph = None
        texts = self._texts(render_docx(content, fmt))
        self.assertIn("Built the thing", texts)
        self.assertNotIn("• Built the thing", texts)

    def test_bullet_indent_applied(self):
        content = _sample_content()
        fmt = default_format_profile()
        fmt.bullet_left_indent_in = 0.5
        blob = render_docx(content, fmt)
        doc = docx.Document(io.BytesIO(blob))
        bullet_para = next(
            p for p in doc.paragraphs if p.text.startswith("• Built")
        )
        self.assertAlmostEqual(
            bullet_para.paragraph_format.left_indent.inches, 0.5
        )

    def test_margins_and_columns_applied(self):
        content = _sample_content()
        fmt = default_format_profile()
        fmt.page_margins_in = {
            "top": 0.5, "right": 0.75, "bottom": 0.5, "left": 0.75,
        }
        fmt.columns = 2
        blob = render_docx(content, fmt)
        doc = docx.Document(io.BytesIO(blob))
        section = doc.sections[0]
        self.assertAlmostEqual(section.top_margin.inches, 0.5)
        self.assertAlmostEqual(section.left_margin.inches, 0.75)
        self.assertAlmostEqual(section.right_margin.inches, 0.75)

    def test_education_description_and_field_render(self):
        content = _sample_content()
        content.education[0].description = "Thesis on parsers"
        texts = "\n".join(
            self._texts(render_docx(content, default_format_profile()))
        )
        self.assertIn("CS", texts)
        self.assertIn("Thesis on parsers", texts)

    def test_contact_bits_skip_empty_fields(self):
        content = _sample_content()
        content.header = ResumeHeader(name="Jane Doe", email="jane@x.com")
        texts = self._texts(render_docx(content, default_format_profile()))
        self.assertIn("jane@x.com", texts)
        # Only the name + one contact line, no separator noise.
        contact_lines = [t for t in texts if "|" in t]
        self.assertEqual(contact_lines, [])


if __name__ == "__main__":
    unittest.main()
