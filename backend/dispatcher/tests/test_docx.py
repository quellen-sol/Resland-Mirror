"""doc_parsers/docx tests — stdlib unittest, builds real .docx in memory.

Run: .venv/bin/python -m unittest backend.dispatcher.tests.test_docx -v
     (from the repo root with the workspace venv)
"""

import io
import unittest

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from dispatcher.doc_parsers.docx import (
    block_sequence,
    column_counts,
    parse_docx,
    raw_text,
    whitespace_map,
)


def _doc_bytes(build) -> bytes:
    doc = Document()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class ParseDocx(unittest.TestCase):
    def test_roundtrips_bytes_to_document(self):
        blob = _doc_bytes(lambda d: d.add_paragraph("hi"))
        doc = parse_docx(blob)
        self.assertEqual(doc.paragraphs[0].text, "hi")


class RawText(unittest.TestCase):
    def test_joins_paragraph_texts(self):
        doc = Document()
        doc.add_paragraph("one")
        doc.add_paragraph("")
        doc.add_paragraph("two")
        self.assertEqual(raw_text(doc), "one\n\ntwo")


class BlockSequence(unittest.TestCase):
    def test_paragraphs_with_styles(self):
        doc = Document()
        doc.add_paragraph("Title here", style="Title")
        doc.add_paragraph("plain")
        blocks = block_sequence(doc)
        self.assertEqual(blocks[0]["type"], "paragraph")
        self.assertEqual(blocks[0]["text"], "Title here")
        self.assertEqual(blocks[0]["style"], "Title")
        self.assertFalse(blocks[0]["ends_section"])
        self.assertEqual(blocks[1]["style"], "Normal")

    def test_table_blocks(self):
        doc = Document()
        doc.add_paragraph("before")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "a1"
        table.cell(0, 1).text = "b1"
        table.cell(1, 0).text = "a2"
        table.cell(1, 1).text = "b2"
        blocks = block_sequence(doc)
        table_block = next(b for b in blocks if b["type"] == "table")
        self.assertEqual(
            table_block["rows"], [["a1", "b1"], ["a2", "b2"]]
        )

    def test_section_break_blocks(self):
        doc = Document()
        doc.add_paragraph("section one")
        doc.add_section()
        doc.add_paragraph("section two")

        blocks = block_sequence(doc)
        types = [b["type"] for b in blocks]
        self.assertIn("section_break", types)
        # The paragraph carrying the break is flagged ends_section.
        ended = [b for b in blocks if b.get("ends_section")]
        self.assertTrue(ended)


class WhitespaceMap(unittest.TestCase):
    def test_empty_and_spacing_flags(self):
        doc = Document()
        p1 = doc.add_paragraph("text")
        p1.paragraph_format.space_before = Pt(12)
        p1.paragraph_format.space_after = Pt(6)
        doc.add_paragraph("   ")

        m = whitespace_map(doc)
        self.assertEqual(len(m), 2)
        self.assertFalse(m[0]["empty"])
        self.assertEqual(m[0]["space_before_pt"], 12.0)
        self.assertEqual(m[0]["space_after_pt"], 6.0)
        self.assertTrue(m[1]["empty"])


class ColumnCounts(unittest.TestCase):
    def test_single_column_default(self):
        doc = Document()
        self.assertEqual(column_counts(doc), [1])

    def test_multi_column_section(self):
        doc = Document()
        section = doc.sections[0]
        cols = section._sectPr.find(qn("w:cols"))
        cols.set(qn("w:num"), "2")
        self.assertEqual(column_counts(doc), [2])


if __name__ == "__main__":
    unittest.main()
