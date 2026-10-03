"""Fidelity tests against real-world PDF resume layouts (issue #49).

fixtures/canva_figures_original.pdf is a Canva export (its text is a
fictional persona) — its name, contact block, section headers, entry
titles, and right-floated dates all live inside pdfminer LTFigures, and its
bullet glyphs are drawn curves, not text. The old extractor dropped all of
that: the agent never saw the profile fields and the captured format came
from a few stray text lines (false two-column layout, 2.25" top margin,
glyph-less bullets, " , " separators).
fixtures/canva_figures_generated_bad.docx is that output; tests pin it as
bad and pin the fixed pipeline's output as good.

Run: .venv/bin/python -m unittest backend.dispatcher.tests.test_pdf_fidelity -v
"""

import io
import unittest
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from shared.docgen import (
    ROLE_ENTRY_TITLE,
    ROLE_LOCATION,
    ROLE_NAME,
    ROLE_SECTION_HEADER,
    SECTION_EDUCATION,
    FormatProfile,
    ResumeContent,
    ResumeEducation,
    ResumeExperience,
    ResumeHeader,
    ResumeSkillGroup,
    capture_format_profile_pdf,
    render_docx,
)

from dispatcher.doc_parsers.pdf import column_counts, parse_pdf, raw_text

try:
    import pdfminer.layout  # noqa: F401

    HAS_PDFMINER = True
except ImportError:
    HAS_PDFMINER = False

FIXTURES = Path(__file__).parent / "fixtures"
ORIGINAL = FIXTURES / "canva_figures_original.pdf"
GENERATED_BAD = FIXTURES / "canva_figures_generated_bad.docx"
FOOD_SERVICE = FIXTURES / "food_service_original.pdf"


# Duck-typed stand-in for the parser's UserDocxProfile — capture only
# reads attributes.


@dataclass
class _Bullet:
    text: str


@dataclass
class _Basic:
    name: str = ""
    email: str = ""
    phone: str = ""
    location: str = ""
    links: list = field(default_factory=list)


@dataclass
class _Exp:
    company: str
    position: str
    start_date: str
    end_date: str | None
    bullets: list = field(default_factory=list)


@dataclass
class _Cat:
    category: str
    skills: list


@dataclass
class _Proj:
    name: str
    link: str | None
    description: str
    start_date: str
    end_date: str | None


@dataclass
class _Edu:
    institution: str
    degree: str
    field_of_study: str | None = None
    start_date: str = ""
    end_date: str | None = None
    description: str | None = None


@dataclass
class _Profile:
    summary: str | None
    basic_info: _Basic
    experience: list = field(default_factory=list)
    skills: list = field(default_factory=list)
    education: list = field(default_factory=list)
    projects: list = field(default_factory=list)


def _canva_profile() -> _Profile:
    """What the parse agent should extract from the pdf fixture."""
    return _Profile(
        summary=(
            "Detail-driven coordinator with a strong background in "
            "operations, event scheduling and logistics. Skilled at building "
            "timelines, aligning stakeholders, and keeping deliverables "
            "moving in busy environments."
        ),
        basic_info=_Basic(
            name="Dana Prescott",
            email="dprescott@example.com",
            phone="(208) 555-0142",
            location="Cedar Hollow Falls, Idaho",
        ),
        experience=[
            _Exp(
                company="Fernhill Creative Studio",
                position="Content Project Manager (Contractor)",
                start_date="Sept 2024",
                end_date="Current",
                bullets=[
                    _Bullet(
                        "Plan weekly posting calendars with realistic "
                        "deadlines and a steady mix of content"
                    ),
                    _Bullet(
                        "Maintained shared asset libraries, caption "
                        "drafts, and reporting sheets"
                    ),
                ],
            ),
            _Exp(
                company="Ember Events",
                position="Owner",
                start_date="January 2022",
                end_date="Current",
                bullets=[_Bullet("Run every part of the business")],
            ),
            _Exp(
                company="Corvia Inc",
                position="Executive Assistant",
                start_date="Dec 2023",
                end_date="Mar 2024",
                bullets=[_Bullet("Kept executive calendars, travel, and action")],
            ),
            _Exp(
                company="Delmar Tax Partners",
                position="Sales Administrative Assistant",
                start_date="Feb 2023",
                end_date="Dec 2023",
                bullets=[_Bullet("Logged outreach with prospects")],
            ),
            _Exp(
                company="Capitol Print Systems",
                position="Administrative Assistant",
                start_date="Nov 2020",
                end_date="Dec 2022",
                bullets=[_Bullet("Processed customer print orders")],
            ),
            _Exp(
                company="Riverside Makers Circle",
                position="Member Coordinator",
                start_date="April 2024",
                end_date="Current",
                bullets=[_Bullet("Take part in monthly peer")],
            ),
        ],
        skills=[
            _Cat(
                "",
                [
                    "Agile & Scrum Fundamentals",
                    "Project Scheduling & Coordination",
                    "Documentation & Record Keeping",
                    "Client & Stakeholder Communication",
                    "Vendor & Cross-Team Collaboration",
                    "Timeline & Task Management",
                    "Budgeting & Resource Allocation",
                    "Process Improvement",
                ],
            )
        ],
        education=[
            _Edu(
                institution="Hartwell Academy",
                degree="Operations Planning Professional Certificate",
                start_date="",
                end_date="December 2024",
            )
        ],
    )


def _canva_content() -> ResumeContent:
    """The same resume as canonical ResumeContent for rendering."""
    return ResumeContent(
        header=ResumeHeader(
            name="Dana Prescott",
            email="dprescott@example.com",
            phone="(208) 555-0142",
            location="Cedar Hollow Falls, Idaho",
        ),
        summary=_canva_profile().summary,
        experiences=[
            ResumeExperience(
                company="Fernhill Creative Studio",
                title="Content Project Manager (Contractor)",
                start_date="Sept 2024",
                end_date="Current",
                bullets=[
                    "Plan weekly posting calendars with realistic deadlines",
                    "Maintained shared asset libraries and caption drafts",
                ],
            ),
            ResumeExperience(
                company="Riverside Makers Circle",
                title="Member Coordinator",
                start_date="April 2024",
                end_date="Current",
                bullets=["Take part in monthly peer mentoring circles"],
            ),
        ],
        skill_groups=[
            ResumeSkillGroup(
                category=None,
                skills=[
                    "Agile & Scrum Fundamentals",
                    "Project Scheduling & Coordination",
                ],
            )
        ],
        education=[
            ResumeEducation(
                institution="Hartwell Academy",
                degree="Operations Planning Professional Certificate",
                field_of_study="",
                start_date="",
                end_date="December 2024",
            )
        ],
        projects=[],
    )


@unittest.skipUnless(HAS_PDFMINER, "pdfminer.six not installed")
class FixtureParse(unittest.TestCase):
    """The extractor must see the text that lives inside LTFigures —
    the regression that emptied this resume's profile fields."""

    @classmethod
    def setUpClass(cls):
        cls.doc = parse_pdf(ORIGINAL.read_bytes())

    def test_name_and_contact_lines_exist(self):
        texts = [line.text for line in self.doc.lines]
        joined = "\n".join(texts)
        self.assertIn("DANA PRESCOTT", joined)
        self.assertIn("dprescott@example.com", joined)
        self.assertIn("(208) 555-0142", joined)
        self.assertIn("Cedar Hollow Falls, Idaho", joined)

    def test_label_value_rows_merge(self):
        self.assertTrue(
            any(
                "Address:" in line.text
                and "Cedar Hollow Falls, Idaho" in line.text
                for line in self.doc.lines
            )
        )

    def test_section_headers_extracted(self):
        texts = {line.text for line in self.doc.lines}
        for title in (
            "SUMMARY",
            "SKILLS",
            "WORK EXPERIENCE",
            "COMMUNITY AND PROFESSIONAL DEVELOPMENT",
        ):
            self.assertIn(title, texts)

    def test_title_and_dates_share_one_line(self):
        line = next(
            l for l in self.doc.lines if "Fernhill Creative Studio" in l.text
        )
        self.assertIn("Sept 2024", line.text)
        self.assertGreaterEqual(line.gap_before_pt, 5.0)

    def test_bullet_dots_become_glyphs(self):
        bullets = [
            line for line in self.doc.lines if line.text.startswith("• ")
        ]
        self.assertGreaterEqual(len(bullets), 8)
        self.assertTrue(
            any("Plan weekly posting calendars" in b.text for b in bullets)
        )

    def test_single_column_both_pages(self):
        self.assertEqual(column_counts(self.doc), [1, 1])

    def test_first_line_is_the_name(self):
        self.assertEqual(self.doc.lines[0].text, "DANA PRESCOTT")

    def test_skills_grid_cells_are_separate_lines(self):
        """The two-band skills grid: every cell its own bulleted line, all
        between SKILLS and WORK EXPERIENCE — no '• A    • B' rows and no
        page-wide column split."""
        texts = [line.text for line in self.doc.lines]
        start = texts.index("SKILLS")
        self.assertEqual(texts[start + 11], "WORK EXPERIENCE")
        cells = texts[start + 1 : start + 11]
        self.assertEqual(
            sorted(cells),
            sorted(
                "• " + s
                for s in (
                    "Agile & Scrum Fundamentals",
                    "Project Scheduling & Coordination",
                    "Documentation & Record Keeping",
                    "Client & Stakeholder Communication",
                    "Vendor & Cross-Team Collaboration",
                    "Timeline & Task Management",
                    "Budgeting & Resource Allocation",
                    "Process Improvement",
                    "Self-Directed & Cross-Functional Leadership",
                    "Content Creation Through Multiple Platforms",
                )
            ),
        )
        # Each band keeps its own top-to-bottom order.
        left = [t for t in cells if "Agile" in t or "Documentation" in t or "Vendor" in t]
        self.assertEqual(left, [c for c in cells if c in left])
        self.assertLess(cells.index(left[0]), cells.index(left[1]))

    def test_contact_labels_keep_their_own_runs(self):
        line = next(l for l in self.doc.lines if l.text.startswith("Phone:"))
        self.assertTrue(line.style_at(0).bold)  # Inter-SemiBold label
        value_at = line.text.index("(208)")
        self.assertFalse(line.style_at(value_at).bold)
        self.assertAlmostEqual(line.run_at(value_at).x0, 228.7, delta=0.5)



@unittest.skipUnless(HAS_PDFMINER, "pdfminer.six not installed")
class FixtureCapture(unittest.TestCase):
    """capture_format_profile_pdf on the real uploaded pdf."""

    @classmethod
    def setUpClass(cls):
        cls.doc = parse_pdf(ORIGINAL.read_bytes())
        cls.fmt = capture_format_profile_pdf(cls.doc, _canva_profile())

    def test_margins_from_full_content(self):
        m = self.fmt.page_margins_in
        self.assertAlmostEqual(m["top"], 0.5, places=1)
        self.assertAlmostEqual(m["left"], 0.8, places=1)
        # The old 2.25" top margin came from dropped header lines.
        self.assertLess(m["top"], 1.0)

    def test_single_column(self):
        self.assertEqual(self.fmt.columns, 1)

    def test_name_and_header_roles(self):
        name = self.fmt.roles[ROLE_NAME]
        self.assertEqual(name.font_name, "Inter")
        self.assertGreaterEqual(name.size_pt, 18.0)
        self.assertTrue(name.bold)
        header = self.fmt.roles[ROLE_SECTION_HEADER]
        self.assertTrue(header.bold)
        self.assertEqual(header.font_name, "Inter")
        self.assertEqual(header.color, "00366D")

    def test_section_titles(self):
        self.assertEqual(self.fmt.section_titles["summary"], "SUMMARY")
        self.assertEqual(self.fmt.section_titles["experience"], "WORK EXPERIENCE")
        self.assertEqual(self.fmt.section_titles["skills"], "SKILLS")
        # A wrapped bullet tail ("professional growth.") must not be
        # claimed as the education header.
        self.assertNotEqual(
            self.fmt.section_titles.get(SECTION_EDUCATION),
            "professional growth.",
        )

    def test_section_order(self):
        self.assertEqual(
            self.fmt.section_order[:3], ["summary", "skills", "experience"]
        )

    def test_entry_line_layout(self):
        # 'Fernhill Creative Studio - Content Project Manager' — the
        # company-first separator is a plain hyphen.
        self.assertEqual(self.fmt.entry_title_separator, " - ")
        self.assertTrue(self.fmt.entry_title_tail_first)
        self.assertTrue(self.fmt.dates_on_title_line)
        self.assertEqual(self.fmt.dates_separator, " – ")
        self.assertEqual(self.fmt.dates_open_label, "Current")
        title = self.fmt.roles[ROLE_ENTRY_TITLE]
        self.assertTrue(title.bold)
        self.assertEqual(title.font_name, "Inter")

    def test_bullet_look(self):
        self.assertEqual(self.fmt.bullet_glyph, "•")
        self.assertIsNotNone(self.fmt.bullet_left_indent_in)

    def test_location_on_own_line(self):
        self.assertTrue(self.fmt.location_on_own_line)
        self.assertEqual(
            self.fmt.roles[ROLE_LOCATION].font_name, "Inter"
        )

    def test_profile_serializes(self):
        self.assertEqual(
            FormatProfile.from_json(self.fmt.to_json()), self.fmt
        )

    def test_sections_keep_their_own_entry_order(self):
        # Experience is company-first; the education line is
        # "Certificate (Hartwell Academy)" — title-first, bracketed.
        self.assertTrue(self.fmt.entry_title_tail_first)
        self.assertEqual(self.fmt.education_title_separator, " (")
        self.assertFalse(self.fmt.education_title_tail_first)

    def test_profile_serializes(self):
        self.assertEqual(
            FormatProfile.from_json(self.fmt.to_json()), self.fmt
        )



@unittest.skipUnless(HAS_PDFMINER, "pdfminer.six not installed")
class FixtureRender(unittest.TestCase):
    """render_docx through the captured profile keeps the source's
    layout and — critically — its header fields."""

    @classmethod
    def setUpClass(cls):
        cls.doc = parse_pdf(ORIGINAL.read_bytes())
        cls.fmt = capture_format_profile_pdf(cls.doc, _canva_profile())
        cls.rendered = Document(io.BytesIO(render_docx(_canva_content(), cls.fmt)))
        cls.paras = cls.rendered.paragraphs
        cls.texts = [p.text for p in cls.paras if p.text.strip()]

    def test_header_fields_present(self):
        joined = "\n".join(self.texts)
        for field_text in (
            "Dana Prescott",
            "dprescott@example.com",
            "(208) 555-0142",
            "Cedar Hollow Falls, Idaho",
        ):
            self.assertIn(field_text, joined)

    def test_section_headers_bold_and_titled(self):
        titles = {
            p.text.strip(): p
            for p in self.paras
            if p.text.strip()
            in ("SUMMARY", "SKILLS", "WORK EXPERIENCE", "EDUCATION")
        }
        for want in ("SUMMARY", "SKILLS", "WORK EXPERIENCE"):
            self.assertIn(want, titles)
            self.assertTrue(titles[want].runs[0].font.bold)

    def test_dates_on_title_line(self):
        title = next(
            p for p in self.paras if "Content Project Manager" in p.text
        )
        self.assertIn("Sept 2024 – Current", title.text)

    def test_source_entry_order_kept(self):
        # Source writes 'Company - Title'; the render must not swap it.
        title = next(
            p for p in self.paras if "Content Project Manager" in p.text
        )
        self.assertTrue(title.text.startswith("Fernhill Creative Studio"))

    def test_comma_separator_has_no_leading_space(self):
        joined = "\n".join(self.texts)
        self.assertNotIn(" , ", joined)

    def test_bullets_rendered(self):
        self.assertTrue(
            any(t.startswith("•") for t in self.texts)
        )

    def test_no_two_column_section(self):
        cols = self.rendered.sections[0]._sectPr.find(qn("w:cols"))
        num = cols.get(qn("w:num")) if cols is not None else None
        self.assertIn(num, (None, "1"))

    def test_education_entry_keeps_brackets_and_drops_dangling_dash(self):
        edu = next(p for p in self.paras if "Hartwell Academy" in p.text)
        self.assertEqual(
            edu.text,
            "Operations Planning Professional Certificate (Hartwell Academy)\tDecember 2024",
        )



def _food_service_profile() -> _Profile:
    """What the parse agent should extract from the food-service pdf —
    a different layout family than the Canva fixture: non-bold name, dates
    on their own line, literal '●' bullets."""
    return _Profile(
        summary=(
            "I am a dependable team player with strong communication, guest "
            "service, and time management skills"
        ),
        basic_info=_Basic(
            name="Dana Prescott",
            email="dprescott@example.com",
            phone="208.555.0142",
        ),
        experience=[
            _Exp(
                company="Corvia Inc",
                position="Executive Assistant",
                start_date="December 2023",
                end_date="Current",
                bullets=[_Bullet("Coordinated calendars, travel, and all meeting")],
            ),
            _Exp(
                company="Chestnut Ramen House",
                position="Team Member",
                start_date="March 2022",
                end_date="May 2022",
                bullets=[_Bullet("Made broths and toppings from scratch")],
            ),
            _Exp(
                company="Sorrel",
                position="Shift Lead",
                start_date="September 2020",
                end_date="August 2021",
                bullets=[_Bullet("Assembled salads, bowls, and sides")],
            ),
        ],
        education=[
            _Edu(
                institution="The Charter School of Merriton",
                degree="High School",
                start_date="",
                end_date="June 2021",
            )
        ],
    )


def _food_service_content() -> ResumeContent:
    return ResumeContent(
        header=ResumeHeader(
            name="Dana Prescott",
            email="dprescott@example.com",
            phone="208.555.0142",
        ),
        summary=_food_service_profile().summary,
        experiences=[
            ResumeExperience(
                company="Corvia Inc",
                title="Executive Assistant",
                start_date="December 2023",
                end_date="Current",
                bullets=["Coordinated calendars, travel, and all meeting logistics"],
            )
        ],
        skill_groups=[],
        education=[
            ResumeEducation(
                institution="The Charter School of Merriton",
                degree="High School",
                field_of_study="",
                start_date="",
                end_date="June 2021",
            )
        ],
        projects=[],
    )


@unittest.skipUnless(HAS_PDFMINER, "pdfminer.six not installed")
class FoodServiceFixture(unittest.TestCase):
    """Second real-pdf fixture — exercises the extractor and capture on a
    layout with literal bullet glyphs, per-line dates, no location, and
    three pages."""

    @classmethod
    def setUpClass(cls):
        cls.doc = parse_pdf(FOOD_SERVICE.read_bytes())
        cls.fmt = capture_format_profile_pdf(cls.doc, _food_service_profile())

    def test_header_fields_parse(self):
        joined = "\n".join(l.text for l in self.doc.lines)
        for field_text in (
            "Dana Prescott",
            "dprescott@example.com",
            "208.555.0142",
        ):
            self.assertIn(field_text, joined)

    def test_three_pages_single_column(self):
        self.assertEqual(len(self.doc.pages), 3)
        self.assertEqual(column_counts(self.doc), [1, 1, 1])

    def test_literal_bullet_glyphs(self):
        bullets = [
            l for l in self.doc.lines if l.text.startswith("●")
        ]
        self.assertGreaterEqual(len(bullets), 10)

    def test_capture_layout(self):
        self.assertEqual(self.fmt.columns, 1)
        self.assertFalse(self.fmt.dates_on_title_line)
        self.assertEqual(self.fmt.bullet_glyph, "●")
        self.assertEqual(
            self.fmt.section_titles["summary"], "ABOUT ME"
        )
        self.assertEqual(
            self.fmt.section_titles["experience"], "EXPERIENCE"
        )
        self.assertEqual(
            self.fmt.section_titles["education"], "EDUCATION"
        )
        self.assertTrue(self.fmt.entry_title_tail_first)
        name = self.fmt.roles[ROLE_NAME]
        self.assertEqual(name.font_name, "ProximaNova")
        self.assertEqual(
            self.fmt.roles[ROLE_SECTION_HEADER].color, "00AB44"
        )

    def test_render_docx_keeps_fields(self):
        rendered = Document(
            io.BytesIO(render_docx(_food_service_content(), self.fmt))
        )
        joined = "\n".join(p.text for p in rendered.paragraphs)
        for field_text in (
            "Dana Prescott",
            "dprescott@example.com",
            "208.555.0142",
            "ABOUT ME",
            "Corvia Inc",
        ):
            self.assertIn(field_text, joined)
        # Source order is 'Company - Title' — the render keeps it.
        title = next(
            p for p in rendered.paragraphs if "Executive Assistant" in p.text
        )
        self.assertTrue(title.text.startswith("Corvia Inc"))

    def test_company_first_bold_attribution(self):
        """'Corvia Inc' is bold dark gray, ' - Executive Assistant' italic
        light gray (ProximaNova-RegularIt), dates 10pt gray on their own
        line — the render keeps each look on its own part."""
        self.assertTrue(self.fmt.entry_title_tail_sub)
        rendered = Document(
            io.BytesIO(render_docx(_food_service_content(), self.fmt))
        )
        paras = rendered.paragraphs
        title = next(p for p in paras if "Executive Assistant" in p.text)
        head, tail = title.runs[0], title.runs[1]
        self.assertEqual(head.text, "Corvia Inc")
        self.assertTrue(head.font.bold)
        self.assertEqual(str(head.font.color.rgb), "353744")
        self.assertEqual(tail.text, " - Executive Assistant")
        self.assertFalse(tail.font.bold)
        self.assertTrue(tail.font.italic)
        self.assertEqual(str(tail.font.color.rgb), "666666")
        dates = next(p for p in paras if p.text == "December 2023 - Current")
        self.assertEqual(dates.runs[0].font.size.pt, 10.0)
        # Education is institution-first in this source.
        edu = next(p for p in paras if "High School" in p.text)
        self.assertEqual(edu.text, "The Charter School of Merriton - High School")
        self.assertTrue(edu.runs[0].font.bold)
        self.assertTrue(edu.runs[1].font.italic)



class GeneratedBadRegression(unittest.TestCase):
    """Pin the fixture's old defects so the assertions above keep meaning:
    the bad output really is missing the header block and really did set
    the page to two columns."""

    @classmethod
    def setUpClass(cls):
        cls.bad = Document(GENERATED_BAD)
        cls.joined = "\n".join(p.text for p in cls.bad.paragraphs)

    def test_bad_fixture_missing_header_fields(self):
        for field_text in (
            "Dana",
            "dprescott@example.com",
            "(208) 555-0142",
            "Cedar Hollow",
        ):
            self.assertNotIn(field_text, self.joined)

    def test_bad_fixture_was_two_column(self):
        cols = self.bad.sections[0]._sectPr.find(qn("w:cols"))
        self.assertIsNotNone(cols)
        self.assertEqual(cols.get(qn("w:num")), "2")

    def test_bad_fixture_had_space_comma(self):
        self.assertIn(" , ", self.joined)


if __name__ == "__main__":
    unittest.main()
