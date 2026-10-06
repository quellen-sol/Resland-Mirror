"""Agent tests — the real compiled LangGraph, driven by a scripted fake
chat model and an in-memory checkpointer (no LLM calls, no Postgres).

`init_chat_model`, `PostgresSaver`, and `get_pg_client` are patched in each
agent module; everything between — the graph, tools, interrupt/resume
plumbing, and usage accounting — is the production code.

Run: .venv/bin/python -m unittest backend.dispatcher.tests.test_agents -v
     (from the repo root with the workspace venv)
"""

import io
import unittest
from unittest.mock import MagicMock, patch

import docx
from langchain.messages import AIMessage
from langgraph.checkpoint.memory import InMemorySaver
from langgraph.errors import GraphRecursionError
from shared.checkpoints import CHECKPOINT_CONN_OPTIONS

from dispatcher.agents import docx_parser, suggestions, tailor, writing


class FakeChatModel:
    """bind_tools() returns self; invoke() pops the next scripted message
    and records every message list it's shown (so tests can read the
    ToolMessages the graph fed back)."""

    def __init__(self, script):
        self.script = list(script)
        self.calls = []
        self.bound_tools = []

    def bind_tools(self, tools):
        self.bound_tools = list(tools)
        return self

    def invoke(self, messages):
        self.calls.append(list(messages))
        nxt = self.script.pop(0)
        return nxt() if callable(nxt) else nxt


def ai(
    content="",
    tool_calls=None,
    in_tok=10,
    out_tok=5,
    cost=0.001,
    gen="gen-1",
    invalid_tool_calls=None,
):
    return AIMessage(
        content=content,
        tool_calls=tool_calls or [],
        invalid_tool_calls=invalid_tool_calls or [],
        usage_metadata={
            "input_tokens": in_tok,
            "output_tokens": out_tok,
            "total_tokens": in_tok + out_tok,
        },
        response_metadata={"cost": cost, "id": gen},
    )


def tc(name, args, tc_id="tc-1", **kw):
    return ai(
        tool_calls=[{"name": name, "args": args, "id": tc_id, "type": "tool_call"}],
        **kw,
    )


def _patch_model(module, script):
    return patch.object(module, "init_chat_model", return_value=FakeChatModel(script))


def _patch_infra(module, saver):
    """PostgresSaver → a real (in-memory) checkpointer; get_pg_client → a
    dummy context manager whose conn is never used."""
    cm = MagicMock()
    cm.__enter__ = MagicMock(return_value=MagicMock())
    cm.__exit__ = MagicMock(return_value=False)
    return (
        patch.object(module, "PostgresSaver", return_value=saver),
        patch.object(module, "get_pg_client", return_value=cm),
    )


def sugg_row(**over):
    base = dict(
        pending_question=None,
        status="PROCESSING",
        result=None,
        model=None,
        user_id="u1",
        company_name="Acme",
        description=None,
        start_date="Jan 2020",
        end_date=None,
        title="Engineer",
        bullet_id=7,
        bullet="Built the thing",
    )
    base.update(over)
    return suggestions.FetchSuggestionRow(**base)


def tailor_profile(**over):
    base = dict(
        header=tailor.TailorHeaderRow(
            full_name="Jane", email=None, phone=None, location=None
        ),
        social_links=[],
        experiences=[
            tailor.TailorExperienceRow(
                experience_id=1,
                company_name="Acme",
                title="Engineer",
                start_date="2020",
                end_date=None,
                description=None,
                bullet_id=10,
                bullet="Built the thing",
                bullet_order=0,
            ),
            tailor.TailorExperienceRow(
                experience_id=2,
                company_name="Beta",
                title="Senior",
                start_date="2022",
                end_date=None,
                description=None,
                bullet_id=20,
                bullet="Shipped the other thing",
                bullet_order=0,
            ),
        ],
        skills=[tailor.TailorSkillRow(name="Python", category=None)],
        education=[],
        projects=[],
    )
    base.update(over)
    return tailor.TailorProfile(**base)


def tailor_resume(**over):
    """A resume that passes every check against tailor_profile(): both
    experiences, each with its bullet."""
    base = {
        "summary": None,
        "experiences": [
            {"experience_id": 1, "bullets": [{"bullet_id": 10, "text": "Built x"}]},
            {"experience_id": 2, "bullets": [{"bullet_id": 20, "text": "Shipped y"}]},
        ],
        "skills": ["Python"],
    }
    base.update(over)
    return base


def tailor_run(**over):
    base = dict(
        id="run-1",
        user_id="u1",
        job_description="We need a Python engineer.",
        status="PROCESSING",
        model=None,
    )
    base.update(over)
    return tailor.FetchTailorRunRow(**base)


def _usage():
    return {
        "input_tokens": 0,
        "output_tokens": 0,
        "generation_ids": [],
        "cost_usd": 0.0,
        "all_costed": True,
    }


def _last_tool_msg(model):
    """The most recent ToolMessage the model was shown."""
    for msg in reversed(model.calls[-1]):
        if getattr(msg, "type", None) == "tool":
            return msg
    return None


def _opts(bullet_id=7, options=("a", "b", "c")):
    """submit/verify args: one bullet with its rewrite options."""
    return {
        "suggestions": {
            "suggestions": [{"bullet_id": bullet_id, "options": list(options)}]
        }
    }


class SuggestionAgent(unittest.TestCase):
    def setUp(self):
        self.saver = InMemorySaver()
        self.saver_patch, self.pg_patch = _patch_infra(suggestions, self.saver)
        self.pg = self.pg_patch.start()
        self.saver_patch.start()
        self.addCleanup(self.pg_patch.stop)
        self.addCleanup(self.saver_patch.stop)

    def _run(
        self,
        script,
        session_id="sess-1",
        answer=None,
        rows=None,
        usage=None,
        prompt=None,
    ):
        with _patch_model(suggestions, script) as model_cls:
            outcome = suggestions.build_and_invoke_agent(
                session_id,
                answer,
                rows if rows is not None else [sugg_row()],
                "openrouter:test/model",
                usage if usage is not None else _usage(),
                prompt,
            )
        return outcome, model_cls.return_value

    def test_submit_waits_for_review_with_options(self):
        outcome, model = self._run(
            [
                tc("get_user_bullet_points", {}),
                tc("verify_final_suggestions", _opts(), tc_id="tc-2"),
                tc("submit_suggestions", _opts(), tc_id="tc-3"),
            ]
        )
        # Submitting pauses on the review interrupt: the options ride along
        # so the runner stores them while the user picks one or asks for more.
        self.assertEqual(outcome.status, "waiting")
        self.assertEqual(outcome.question, suggestions.REVIEW_PROMPT)
        self.assertEqual(outcome.result.suggestions[0].options, ["a", "b", "c"])
        # The model isn't called again after a successful submit.
        self.assertEqual(len(model.calls), 3)

    def test_reply_to_review_resumes_and_resubmits(self):
        self._run([tc("submit_suggestions", _opts())])
        outcome, model = self._run(
            [tc("submit_suggestions", _opts(options=("d", "e", "f")))],
            answer="shorter please",
        )
        self.assertEqual(outcome.status, "waiting")
        self.assertEqual(outcome.result.suggestions[0].options, ["d", "e", "f"])
        # The reply reaches the model as a user message with the earlier
        # options still in context, so it can avoid repeating them.
        reply = model.calls[0][-1]
        self.assertEqual(reply.type, "human")
        self.assertEqual(reply.content, "shorter please")
        earlier = [
            m
            for m in model.calls[0]
            if getattr(m, "type", None) == "ai" and m.tool_calls
        ]
        self.assertEqual(earlier[0].tool_calls[0]["name"], "submit_suggestions")

    def test_reply_without_resubmit_completes_with_no_new_result(self):
        self._run([tc("submit_suggestions", _opts())])
        outcome, _ = self._run([ai("Glad they work for you.")], answer="thanks")
        # No fresh submit: the runner keeps the stored options and completes.
        self.assertEqual(outcome.status, "completed")
        self.assertIsNone(outcome.result)

    def test_bullet_points_tool_returns_experience_rows(self):
        _, model = self._run(
            [
                tc("get_user_bullet_points", {}),
                tc("submit_suggestions", _opts(), tc_id="tc-2"),
            ]
        )
        tool_msg = model.calls[1][-1]
        self.assertIn('"Acme"', tool_msg.content)
        self.assertIn('"Built the thing"', tool_msg.content)
        # Session-internal fields are stripped from what the model sees.
        self.assertNotIn("user_id", tool_msg.content)
        self.assertNotIn("pending_question", tool_msg.content)

    def test_clarifying_question_waits_with_question(self):
        outcome, _ = self._run(
            [tc("ask_clarifying_question", {"question": "How many engineers?"})]
        )
        self.assertEqual(outcome.status, "waiting")
        self.assertEqual(outcome.question, "How many engineers?")
        self.assertIsNone(outcome.result)

    def test_answer_resumes_interrupt(self):
        # First run ends waiting on a question; the answer resumes the
        # same checkpoint (shared saver) and the tool returns it.
        self._run([tc("ask_clarifying_question", {"question": "Team size?"})])
        outcome, model = self._run(
            [tc("submit_suggestions", _opts())],
            answer="5 engineers",
        )
        self.assertEqual(outcome.status, "waiting")
        self.assertIsNotNone(outcome.result)
        resumed_msg = model.calls[0][-1]
        self.assertEqual(resumed_msg.content, "5 engineers")

    def test_custom_prompt_seeds_kickoff_message(self):
        outcome, model = self._run(
            [ai("done")],
            prompt="Keep them under 10 words",
        )
        self.assertEqual(outcome.status, "completed")
        kickoff = model.calls[0][-1]
        self.assertEqual(kickoff.type, "human")
        self.assertIn("Keep them under 10 words", kickoff.content)

    def test_retry_resumes_checkpoint_without_new_kickoff(self):
        # First run crashes mid-node — the checkpoint keeps llm_call pending.
        def boom():
            raise RuntimeError("provider down")

        with self.assertRaises(RuntimeError):
            self._run([boom])

        # Retry with no answer: the pending node resumes in place instead of
        # a duplicate kickoff message being injected.
        outcome, model = self._run([tc("submit_suggestions", _opts())])
        self.assertEqual(outcome.status, "waiting")
        humans = [
            m
            for m in model.calls[0]
            if getattr(m, "type", None) == "human"
        ]
        self.assertEqual(len(humans), 1)
        self.assertIn("get_user_bullet_points", humans[0].content)

    def test_checkpointer_connects_to_its_schema(self):
        self._run([ai("done")])
        self.assertEqual(
            self.pg.call_args.kwargs["options"], CHECKPOINT_CONN_OPTIONS
        )

    def test_answer_on_finished_run_hits_empty_input_fallback(self):
        # Command(resume) on a checkpoint with no pending interrupt raises
        # EmptyInputError — the code falls back to invoke(None).
        self._run([ai("done")])
        outcome, _ = self._run([], answer="late answer")
        self.assertEqual(outcome.status, "completed")

    def test_usage_accumulates_tokens_cost_and_generation_ids(self):
        usage = _usage()
        self._run(
            [
                tc("get_user_bullet_points", {}, in_tok=10, out_tok=5, cost=0.001),
                tc(
                    "submit_suggestions",
                    _opts(),
                    in_tok=20,
                    out_tok=8,
                    cost=0.002,
                    gen="gen-2",
                    tc_id="tc-2",
                ),
            ],
            usage=usage,
        )
        self.assertEqual(usage["input_tokens"], 30)
        self.assertEqual(usage["output_tokens"], 13)
        self.assertAlmostEqual(usage["cost_usd"], 0.003)
        self.assertEqual(usage["generation_ids"], ["gen-1", "gen-2"])
        self.assertTrue(usage["all_costed"])

    def test_missing_inline_cost_marks_not_all_costed(self):
        usage = _usage()
        msg = tc("get_user_bullet_points", {})
        msg.response_metadata = {"id": "gen-9"}  # no cost field
        self._run(
            [
                msg,
                tc("submit_suggestions", _opts(), tc_id="tc-2", gen="gen-2"),
            ],
            usage=usage,
        )
        self.assertFalse(usage["all_costed"])
        self.assertEqual(usage["generation_ids"], ["gen-9", "gen-2"])

    def _verify_feedback(self, args):
        _, model = self._run([tc("verify_final_suggestions", args), ai("done")])
        return _last_tool_msg(model).content

    def test_verify_rejects_unknown_bullet_id(self):
        msg = self._verify_feedback(_opts(bullet_id=999))
        self.assertIn("Invalid bullet id (999)", msg)

    def test_verify_requires_a_bullet(self):
        msg = self._verify_feedback({"suggestions": {"suggestions": []}})
        self.assertIn("at least one bullet", msg)

    def test_verify_rejects_duplicate_bullet(self):
        entry = {"bullet_id": 7, "options": ["a", "b", "c"]}
        msg = self._verify_feedback({"suggestions": {"suggestions": [entry, entry]}})
        self.assertIn("Duplicate bullet id (7)", msg)

    def test_verify_requires_three_options(self):
        msg = self._verify_feedback(_opts(options=("a", "b")))
        self.assertIn("exactly 3 options", msg)

    def test_verify_rejects_blank_option(self):
        msg = self._verify_feedback(_opts(options=("a", "  ", "c")))
        self.assertIn("blank option", msg)

    def test_verify_rejects_repeated_options(self):
        msg = self._verify_feedback(_opts(options=("Built x.", "built x", "c")))
        self.assertIn("distinct", msg)

    def test_verify_accepts_valid_options(self):
        msg = self._verify_feedback(_opts())
        self.assertIn("valid", msg)
        self.assertNotIn("Invalid", msg)

    def test_system_prompt_carries_shared_writing_rules(self):
        _, model = self._run([ai("done")])
        system = model.calls[0][0]
        self.assertEqual(system.type, "system")
        self.assertIn(writing.BULLET_WRITING_RULES, system.content)

    def test_submit_rejects_invalid_options_and_keeps_going(self):
        # A bad submit is bounced back to the model instead of being stored;
        # the corrected resubmit is what reaches the user.
        outcome, model = self._run(
            [
                tc("submit_suggestions", _opts(options=("only one",))),
                tc("submit_suggestions", _opts(), tc_id="tc-2"),
            ]
        )
        bounced = model.calls[1][-1]
        self.assertIn("Not submitted", bounced.content)
        self.assertIn("exactly 3 options", bounced.content)
        self.assertEqual(outcome.status, "waiting")
        self.assertEqual(outcome.result.suggestions[0].options, ["a", "b", "c"])


class TailorAgent(unittest.TestCase):
    def setUp(self):
        self.saver = InMemorySaver()
        self.saver_patch, self.pg_patch = _patch_infra(tailor, self.saver)
        self.pg = self.pg_patch.start()
        self.saver_patch.start()
        self.addCleanup(self.pg_patch.stop)
        self.addCleanup(self.saver_patch.stop)

    def _run(self, script, run_id="run-1", answer=None, run=None, profile=None, usage=None):
        with _patch_model(tailor, script) as model_cls:
            outcome = tailor.build_and_invoke_agent(
                run_id,
                answer,
                run if run is not None else tailor_run(),
                profile if profile is not None else tailor_profile(),
                "openrouter:test/model",
                usage if usage is not None else _usage(),
            )
        return outcome, model_cls.return_value

    def _verify_call(self, resume):
        return tc("verify_tailored_resume", {"resume": resume})

    def _ask(self, question="Want a summary?", **kw):
        """A clarifying question — pauses the run so a test can stop there."""
        return tc("ask_clarifying_question", {"question": question}, **kw)

    def test_submit_waits_for_feedback_then_done_completes(self):
        resume = tailor_resume()
        outcome, _ = self._run(
            [
                tc("get_job_description", {}),
                tc("get_user_profile", {}),
                self._verify_call(resume),
                tc("submit_tailored_resume", {"resume": resume}, tc_id="tc-4"),
            ]
        )
        # A submitted resume pauses on the review interrupt — the result
        # rides along so the runner can store it while waiting.
        self.assertEqual(outcome.status, "waiting")
        self.assertEqual(outcome.result.skills, ["Python"])
        self.assertEqual(outcome.result.experiences[0].experience_id, 1)

        # Revision feedback resumes the checkpointed graph — the agent sees
        # the reply as a user message and resubmits, pausing again.
        revised = tailor_resume(
            experiences=[
                {"experience_id": 1, "bullets": [{"bullet_id": 10, "text": "Shipped x"}]},
                {"experience_id": 2, "bullets": [{"bullet_id": 20, "text": "Shipped y"}]},
            ]
        )
        outcome, model = self._run(
            [
                self._verify_call(revised),
                tc("submit_tailored_resume", {"resume": revised}),
            ],
            answer="stronger verbs",
        )
        self.assertEqual(outcome.status, "waiting")
        self.assertEqual(outcome.result.experiences[0].bullets[0].text, "Shipped x")
        self.assertEqual(model.calls[0][-1].content, "stronger verbs")

        # "done" lets the agent end the conversation — no fresh submit, so
        # the outcome carries no result and the runner completes the run.
        outcome, model = self._run([ai("All set!")], answer="done")
        self.assertEqual(outcome.status, "completed")
        self.assertIsNone(outcome.result)
        self.assertEqual(model.calls[0][-1].content, "done")

    def test_system_prompt_carries_shared_writing_rules(self):
        _, model = self._run([self._ask()])
        system = model.calls[0][0]
        self.assertEqual(system.type, "system")
        self.assertIn(writing.BULLET_WRITING_RULES, system.content)

    def test_checkpointer_connects_to_its_schema(self):
        self._run([self._ask()])
        self.assertEqual(
            self.pg.call_args.kwargs["options"], CHECKPOINT_CONN_OPTIONS
        )

    def test_text_reply_before_submit_is_nudged_back_to_tools(self):
        # The shape of a real failed run: both get_* tools in one turn, then
        # the interview questions written as plain text (finish_reason
        # "stop"). They never reach the user, so the graph sends the model
        # back to its tools instead of ending the run.
        fetch_both = ai(
            "I'll start by fetching the job description and your profile.",
            tool_calls=[
                {"name": name, "args": {}, "id": f"tc-{name}", "type": "tool_call"}
                for name in ("get_job_description", "get_user_profile")
            ],
        )
        outcome, model = self._run(
            [
                fetch_both,
                ai("1. How many users?\n2. Want a summary?"),
                self._ask("1. How many users?\n2. Want a summary?"),
            ]
        )
        self.assertEqual(outcome.status, "waiting")
        self.assertEqual(outcome.question, "1. How many users?\n2. Want a summary?")
        nudge = model.calls[2][-1]
        self.assertEqual(nudge.type, "human")
        self.assertEqual(nudge.content, tailor.NUDGE_MESSAGE)

    def test_stalled_model_gives_up_after_max_nudges(self):
        outcome, model = self._run(
            [ai("Let me think.") for _ in range(tailor.MAX_NUDGES + 1)]
        )
        # The runner turns a completed outcome with no result into a failure.
        self.assertEqual(outcome.status, "completed")
        self.assertIsNone(outcome.result)
        self.assertEqual(len(model.calls), tailor.MAX_NUDGES + 1)

    def test_nudge_cap_resets_after_a_tool_call(self):
        stall = [ai("Let me think.") for _ in range(tailor.MAX_NUDGES)]
        outcome, _ = self._run(
            [*stall, tc("get_job_description", {}), *stall, self._ask()]
        )
        self.assertEqual(outcome.status, "waiting")

    def _broken_call(self, name="ask_clarifying_question", tc_id="bad-1"):
        return ai(
            invalid_tool_calls=[
                {
                    "type": "invalid_tool_call",
                    "name": name,
                    "args": '{"question": "1. How many',
                    "id": tc_id,
                    "error": "Unterminated string",
                }
            ]
        )

    def test_malformed_tool_call_gets_its_parse_error_back(self):
        outcome, model = self._run([self._broken_call(), self._ask()])
        self.assertEqual(outcome.status, "waiting")
        # Every tool call id needs an answer, or the provider rejects the
        # next request — the parse error doubles as the retry hint.
        reply = model.calls[1][-1]
        self.assertEqual(reply.type, "tool")
        self.assertEqual(reply.tool_call_id, "bad-1")
        self.assertEqual(reply.status, "error")
        self.assertIn("Unterminated string", reply.content)

    def test_repeated_malformed_calls_give_up_after_max_nudges(self):
        # A model that keeps truncating its submit JSON would otherwise loop
        # (and bill) until LangGraph's recursion limit.
        outcome, model = self._run(
            [
                self._broken_call("submit_tailored_resume", tc_id=f"bad-{i}")
                for i in range(tailor.MAX_NUDGES + 1)
            ]
        )
        self.assertEqual(outcome.status, "completed")
        self.assertIsNone(outcome.result)
        self.assertEqual(len(model.calls), tailor.MAX_NUDGES + 1)

    def test_malformed_calls_and_text_replies_share_one_cap(self):
        outcome, model = self._run(
            [self._broken_call(), ai("Let me think."), ai("Still thinking.")]
        )
        self.assertEqual(outcome.status, "completed")
        self.assertEqual(len(model.calls), tailor.MAX_NUDGES + 1)

    def test_retry_resumes_checkpoint_without_new_kickoff(self):
        # A redelivered first turn — the crashed run left llm_call pending.
        def boom():
            raise RuntimeError("provider down")

        with self.assertRaises(RuntimeError):
            self._run([tc("get_job_description", {}), boom])

        outcome, model = self._run([self._ask()])
        self.assertEqual(outcome.status, "waiting")
        humans = [m for m in model.calls[0] if getattr(m, "type", None) == "human"]
        self.assertEqual(len(humans), 1)
        self.assertIn("get_job_description", humans[0].content)

    def test_redelivered_turn_at_a_question_asks_it_again(self):
        # The turn finished but its ack never landed — the redelivery re-raises
        # the pending question instead of dropping it for a second kickoff.
        self._run([self._ask("1. How many users?")])
        outcome, model = self._run([])
        self.assertEqual(outcome.status, "waiting")
        self.assertEqual(outcome.question, "1. How many users?")
        self.assertEqual(model.calls, [])

    def test_tools_hand_back_job_description_and_profile(self):
        _, model = self._run(
            [
                tc("get_job_description", {}),
                tc("get_user_profile", {}),
                self._ask(),
            ]
        )
        self.assertIn("Python engineer", model.calls[1][-1].content)
        self.assertIn('"Acme"', model.calls[2][-1].content)

    def test_clarifying_question_waits_then_resumes(self):
        self._run([tc("ask_clarifying_question", {"question": "Want a summary?"})])
        run = tailor_run()
        profile = tailor_profile()
        resume = tailor_resume(summary="Yes please", skills=[])
        outcome, model = self._run(
            [tc("submit_tailored_resume", {"resume": resume})],
            answer="yes, write one",
            run=run,
            profile=profile,
        )
        # Submit now pauses on the review interrupt instead of ending.
        self.assertEqual(outcome.status, "waiting")
        self.assertIsNotNone(outcome.result)
        self.assertEqual(model.calls[0][-1].content, "yes, write one")

    def test_usage_accumulates(self):
        usage = _usage()
        resume = tailor_resume(skills=[])
        self._run(
            [
                tc("get_job_description", {}, in_tok=3, out_tok=2, cost=0.0005),
                tc("submit_tailored_resume", {"resume": resume}, in_tok=7, out_tok=4, cost=0.0005, gen="gen-2", tc_id="tc-2"),
            ],
            usage=usage,
        )
        self.assertEqual(usage["input_tokens"], 10)
        self.assertEqual(usage["output_tokens"], 6)
        self.assertAlmostEqual(usage["cost_usd"], 0.001)
        self.assertEqual(usage["generation_ids"], ["gen-1", "gen-2"])

    def _verify_feedback(self, resume):
        _, model = self._run([self._verify_call(resume), self._ask()])
        return _last_tool_msg(model).content

    def test_verify_requires_every_experience(self):
        # Dropping a role leaves a gap in the user's work history — the
        # check names what's missing so the model can put it back.
        only_acme = tailor_resume(
            experiences=[
                {"experience_id": 1, "bullets": [{"bullet_id": 10, "text": "x"}]}
            ]
        )
        msg = self._verify_feedback(only_acme)
        self.assertIn("Every experience must stay", msg)
        self.assertIn("2 (Senior at Beta)", msg)
        self.assertNotIn("1 (Engineer at Acme)", msg)

    def test_verify_lists_every_missing_experience(self):
        msg = self._verify_feedback(tailor_resume(experiences=[]))
        self.assertIn("1 (Engineer at Acme), 2 (Senior at Beta)", msg)

    def test_verify_rejects_an_experience_left_without_bullets(self):
        msg = self._verify_feedback(
            tailor_resume(
                experiences=[
                    {"experience_id": 1, "bullets": [{"bullet_id": 10, "text": "x"}]},
                    {"experience_id": 2, "bullets": []},
                ]
            )
        )
        self.assertIn("2 (Senior at Beta) has no bullets", msg)

    def test_verify_rejects_a_duplicate_experience(self):
        resume = tailor_resume()
        resume["experiences"].append(
            {"experience_id": 1, "bullets": [{"bullet_id": 10, "text": "again"}]}
        )
        msg = self._verify_feedback(resume)
        self.assertIn("Duplicate experience id (1)", msg)

    def test_verify_still_allows_dropping_bullets(self):
        profile = tailor_profile()
        profile.experiences.append(
            profile.experiences[0].model_copy(
                update={"bullet_id": 11, "bullet": "Fixed the printer", "bullet_order": 1}
            )
        )
        _, model = self._run(
            [self._verify_call(tailor_resume()), self._ask()], profile=profile
        )
        self.assertIn("All ids are valid", _last_tool_msg(model).content)

    def test_verify_allows_an_experience_with_no_bullets_to_stay_empty(self):
        profile = tailor_profile()
        profile.experiences.append(
            tailor.TailorExperienceRow(
                experience_id=3,
                company_name="Gamma",
                title="Intern",
                start_date="2018",
                end_date="2019",
                description=None,
                bullet_id=None,
                bullet=None,
                bullet_order=None,
            )
        )
        resume = tailor_resume()
        resume["experiences"].append({"experience_id": 3, "bullets": []})
        _, model = self._run([self._verify_call(resume), self._ask()], profile=profile)
        self.assertIn("All ids are valid", _last_tool_msg(model).content)

    def test_submit_bounces_a_resume_that_drops_an_experience(self):
        # A model can skip verify — submit runs the same check, so a resume
        # missing a role is never stored.
        only_acme = tailor_resume(
            experiences=[
                {"experience_id": 1, "bullets": [{"bullet_id": 10, "text": "x"}]}
            ]
        )
        outcome, model = self._run(
            [
                tc("submit_tailored_resume", {"resume": only_acme}),
                tc("submit_tailored_resume", {"resume": tailor_resume()}, tc_id="tc-2"),
            ]
        )
        bounced = model.calls[1][-1]
        self.assertIn("Not submitted", bounced.content)
        self.assertIn("2 (Senior at Beta)", bounced.content)
        self.assertEqual(outcome.status, "waiting")
        self.assertEqual(
            [e.experience_id for e in outcome.result.experiences], [1, 2]
        )

    def test_resubmit_loop_stops_at_the_recursion_limit(self):
        # Every bounced submit is a valid tool call, so the nudge cap never
        # trips — the per-invocation step limit is what stops the billing.
        # The runner turns the error into a resumable question.
        only_acme = tailor_resume(
            experiences=[
                {"experience_id": 1, "bullets": [{"bullet_id": 10, "text": "x"}]}
            ]
        )
        script = [
            tc("submit_tailored_resume", {"resume": only_acme}, tc_id=f"tc-{i}")
            for i in range(200)
        ]
        with (
            _patch_model(tailor, script) as model_cls,
            self.assertRaises(GraphRecursionError),
        ):
            tailor.build_and_invoke_agent(
                "run-1", None, tailor_run(), tailor_profile(),
                "openrouter:test/model", _usage(),
            )
        calls = len(model_cls.return_value.calls)
        self.assertLessEqual(calls, tailor.RECURSION_LIMIT // 2 + 1)
        self.assertLess(calls, 200)

    def test_verify_rejects_unknown_experience(self):
        msg = self._verify_feedback(
            {
                "summary": None,
                "experiences": [{"experience_id": 99, "bullets": []}],
                "skills": [],
            }
        )
        self.assertIn("Invalid experience id (99)", msg)

    def test_verify_rejects_bullet_under_wrong_experience(self):
        msg = self._verify_feedback(
            {
                "summary": None,
                "experiences": [
                    {"experience_id": 1, "bullets": [{"bullet_id": 20, "text": "x"}]}
                ],
                "skills": [],
            }
        )
        self.assertIn("does not belong", msg)

    def test_verify_rejects_duplicate_bullets(self):
        msg = self._verify_feedback(
            {
                "summary": None,
                "experiences": [
                    {
                        "experience_id": 1,
                        "bullets": [
                            {"bullet_id": 10, "text": "x"},
                            {"bullet_id": 10, "text": "y"},
                        ],
                    }
                ],
                "skills": [],
            }
        )
        self.assertIn("Duplicate bullet id (10)", msg)

    def test_verify_rejects_unknown_skills(self):
        msg = self._verify_feedback(tailor_resume(skills=["Kubernetes"]))
        self.assertIn("Skills not present", msg)

    def test_verify_accepts_valid_resume(self):
        msg = self._verify_feedback(tailor_resume())
        self.assertIn("All ids are valid", msg)

    def test_missing_inline_cost_marks_not_all_costed(self):
        usage = _usage()
        msg = tc("get_user_profile", {})
        msg.response_metadata = {"id": "gen-9"}  # no cost field
        _, _ = self._run(
            [
                msg,
                self._verify_call(tailor_resume()),
                tc("submit_tailored_resume", {"resume": tailor_resume()}, tc_id="tc-3"),
            ],
            usage=usage,
        )
        self.assertFalse(usage["all_costed"])
        self.assertIn("gen-9", usage["generation_ids"])


def _doc():
    buf = io.BytesIO()
    d = docx.Document()
    d.add_paragraph("Jane Doe")
    d.add_paragraph("Engineer at Acme")
    d.save(buf)
    buf.seek(0)
    return docx.Document(buf)


class DocxParserAgent(unittest.TestCase):
    def _build(self):
        """build_resume_agent with create_agent stubbed — returns the
        captured tools/model plus the profile queue."""
        agent = MagicMock()
        with patch.object(
            docx_parser, "create_agent", return_value=agent
        ) as create:
            out = docx_parser.build_resume_agent(_doc())
        _, kwargs = create.call_args
        tools = {t.name: t for t in kwargs["tools"]}
        return out, agent, kwargs, tools

    def test_create_agent_wired_with_model_tools_prompt(self):
        _, _, kwargs, tools = self._build()
        self.assertEqual(kwargs["model"], docx_parser.RESUME_PARSER_MODEL)
        self.assertEqual(kwargs["system_prompt"], docx_parser.DOCX_AGENT_PARSER_SYSTEM_PROMPT)
        self.assertEqual(
            set(tools),
            {
                "get_raw_doc_text",
                "get_block_sequence",
                "get_whitespace_map",
                "get_columns",
                "submit_user_profile",
            },
        )

    def test_tools_return_precomputed_doc_views(self):
        _, _, _, tools = self._build()
        self.assertIn("Jane Doe", tools["get_raw_doc_text"].invoke({}))
        blocks = tools["get_block_sequence"].invoke({})
        self.assertTrue(any("Jane Doe" in str(b) for b in blocks))
        self.assertIsInstance(tools["get_whitespace_map"].invoke({}), list)
        self.assertEqual(tools["get_columns"].invoke({}), [1])

    def test_submit_tool_fills_profile_queue(self):
        out, _, _, tools = self._build()
        profile = docx_parser.UserDocxProfile(
            summary="s",
            basic_info=docx_parser.UserProfileBasicInfo(
                name="Jane", email="", phone="", location=""
            ),
            experience=[],
            skills=[],
            education=[],
            projects=[],
        )
        tools["submit_user_profile"].invoke({"profile": profile})
        self.assertIs(out.profile_queue.get_nowait(), profile)

    def test_agent_parse_docx_invokes_and_returns_queue(self):
        out, agent, _, _ = self._build()
        with patch.object(
            docx_parser, "build_resume_agent", return_value=out
        ):
            q = docx_parser.agent_parse_docx(MagicMock())
        self.assertIs(q, out.profile_queue)
        agent.invoke.assert_called_once_with(
            {"messages": [{"role": "user", "content": "Parse this document"}]}
        )


if __name__ == "__main__":
    unittest.main()
