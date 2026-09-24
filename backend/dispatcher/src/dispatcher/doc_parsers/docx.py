import io
from typing import Any

from docx import Document
from docx.document import Document as _Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def parse_docx(file_data: bytes) -> _Document:
    return Document(io.BytesIO(file_data))


def raw_text(doc: _Document) -> str:
    return "\n".join([paragraph.text for paragraph in doc.paragraphs])


def block_sequence(doc: _Document) -> list[dict[str, Any]]:
    """Document blocks in order: paragraphs (with style name and
    "ends_section" when the paragraph carries a section break), tables
    (as row/cell text), and section breaks."""
    blocks = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, doc)
            pPr = child.find(qn("w:pPr"))
            ends_section = pPr is not None and pPr.find(qn("w:sectPr")) is not None
            blocks.append(
                {
                    "type": "paragraph",
                    "text": para.text,
                    "style": para.style.name if para.style else None,
                    "ends_section": ends_section,
                }
            )
        elif child.tag == qn("w:tbl"):
            table = Table(child, doc)
            blocks.append(
                {
                    "type": "table",
                    "rows": [[cell.text for cell in row.cells] for row in table.rows],
                }
            )
        elif child.tag == qn("w:sectPr"):
            blocks.append({"type": "section_break"})
    return blocks


def whitespace_map(doc: _Document) -> list[dict[str, Any]]:
    """Per paragraph index: whether it's empty, plus its space before/after
    in points. Blank paragraphs and large spacing mark the visual boundaries
    between resume sections."""

    def _pt(value) -> float:
        return round(value.pt, 1) if value is not None else 0.0

    return [
        {
            "index": i,
            "empty": not p.text.strip(),
            "space_before_pt": _pt(p.paragraph_format.space_before),
            "space_after_pt": _pt(p.paragraph_format.space_after),
        }
        for i, p in enumerate(doc.paragraphs)
    ]


def column_counts(doc: _Document) -> list[int]:
    """Number of text columns per document section. A value greater than 1
    means the section uses a multi-column layout."""
    counts = []
    for section in doc.sections:
        cols = section._sectPr.find(qn("w:cols"))
        num = cols.get(qn("w:num")) if cols is not None else None
        counts.append(int(num) if num else 1)
    return counts
